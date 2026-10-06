"""El inventario de actas en Excel: una hoja, una fila por PDF, con presentación ejecutiva.

No es el FUID ni ningún formato oficial: es la tabla que el operador pidió para
recorrer una carpeta de actas con el Excel al lado. Cada fila dice qué archivo
es, qué acta y de qué fecha, el folio inicial y el final, cuántos folios tiene y
un resumen de todo el documento.

Se presenta como un informe y no como un volcado de datos, porque es lo que se
entrega: una banda de título, una fila de indicadores -- cuántas actas, cuántos
folios, qué periodo cubren, cuántas hay que revisar --, la tabla con filtros y
una nota que explica de dónde sale cada dato. Queda lista para imprimir en
horizontal, con el encabezado de la tabla repetido en cada hoja.

Se escribe desde cero con openpyxl, sin plantilla: abrirlo y guardarlo es
instantáneo, y la carpeta lo puede reescribir después de cada acta sin que se
note en el tiempo de la corrida.
"""

from __future__ import annotations

import math
import re
from collections.abc import Mapping, Sequence
from dataclasses import asdict, is_dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..domain.fechas import MESES

#: Cómo se llama la planilla que deja cada acta junto a su trabajo. Es el mismo
#: sufijo que la planilla de un documento, para que el botón «Inventario en
#: Excel» de la pantalla -- que la busca por él -- la encuentre sin cambiar.
SUFIJO = "__inventario.xlsx"

#: El nombre del programa, tal como lo llama el operador, para el pie de página
#: y las propiedades del libro.
PROGRAMA = "División de documentos AFINIA"

#: Las columnas, en el orden en que se leen: qué archivo, qué acta, de cuándo,
#: sus folios y de qué trata. El ancho está pensado para que la tabla se lea en
#: una pantalla, salvo el resumen, que se ajusta dentro de la celda.
COLUMNAS: tuple[tuple[str, str, int], ...] = (
    ("N°", "orden", 6),
    ("Carpeta", "carpeta", 15),
    ("Archivo PDF", "archivo", 30),
    # Qué se hizo con ella: el PDF en que quedó al dividir el libro, o que ya
    # era un acta y sólo se inventarió.
    ("Resultado", "resultado", 34),
    ("Páginas", "paginas", 9),
    ("Acta", "acta", 26),
    ("Fecha de la sesión", "fecha", 28),
    ("Folio inicial", "folio_inicial", 9),
    ("Folio final", "folio_final", 9),
    ("N° de folios", "folios", 9),
    ("Resumen del documento", "resumen", 72),
    # Qué trae cada sesión y en qué páginas: citación, asistencia, acta y
    # anexos. Es la columna que lleva de la fila al papel sin abrir el libro.
    ("Piezas de la sesión", "piezas", 46),
    ("Observaciones", "observaciones", 34),
)

#: Las que llevan un número corto y van centradas.
_CENTRADAS = {"orden", "paginas", "folio_inicial", "folio_final", "folios"}

#: Dónde va cada parte de la hoja.
FILA_DEL_TITULO = 1
FILA_DEL_SUBTITULO = 2
FILA_DE_LA_LINEA = 3
FILA_DE_INDICADORES = 5
FILA_DE_ENCABEZADOS = 8

#: Los indicadores de la cabecera: qué dicen y qué columnas ocupa cada uno.
_INDICADORES = (
    ("ACTAS", 1, 2),
    ("FOLIOS", 3, 3),
    ("PERIODO", 4, 5),
    ("POR REVISAR", 6, 6),
    ("CARPETA", 7, 10),
    ("GENERADO", 11, 11),
)

# La paleta: azul marino institucional, un dorado para la línea del título,
# grises para lo secundario y un ámbar para lo que hay que revisar.
_MARINO = "1F3864"
_AZUL_CLARO = "D9E1F2"
_TARJETA = "F2F5FA"
_DORADO = "C9A227"
_GRIS = "595959"
_GRIS_CLARO = "8C8C8C"
_AMBAR = "9C5700"
_AMBAR_FONDO = "FFF4E5"
_FUENTE = "Calibri"

_LINEA = Side(style="thin", color="D9D9D9")
_CAJA = Border(left=_LINEA, right=_LINEA, top=_LINEA, bottom=_LINEA)

#: Alto de un renglón de texto a tamaño 10, en puntos.
_ALTO_DE_RENGLON = 13.5

_FECHA_EN_CIFRAS = re.compile(r"(\d{1,2})-(\d{1,2})-(\d{4})")
_FECHA_LARGA = re.compile(r"(\d{1,2})\s+de\s+([a-záéíóúñ]+)\s+de\s+(\d{4})", re.IGNORECASE)


def escribir_inventario_de_actas(
    filas: Sequence[Mapping[str, object] | object],
    destino: Path,
    *,
    carpeta: str | None = None,
    generado: datetime | None = None,
    nota: str | None = None,
) -> Path:
    """Escribe el inventario en ``destino`` y devuelve dónde quedó.

    Las filas se numeran aquí, en el orden en que llegan: el de las rutas de los
    PDF, que es el de la carpeta. ``carpeta`` es lo que se dice en la cabecera;
    sin ella se nombra el archivo, que es lo que hay cuando el inventario es de
    un solo PDF.
    """
    datos = [_como_dict(fila) for fila in filas]
    for numero, fila in enumerate(datos, start=1):
        fila["orden"] = numero
        fila["paginas"] = _paginas(fila)
        fila["carpeta"] = fila.get("carpeta") or "—"
    generado = generado or datetime.now()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Inventario de actas"
    hoja.sheet_view.showGridLines = False
    libro.properties.title = "Inventario de actas de consejo"
    libro.properties.creator = PROGRAMA

    for columna, (_titulo, _clave, ancho) in enumerate(COLUMNAS, start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho

    _banda_del_titulo(hoja, carpeta, datos, nota)
    _indicadores(hoja, datos, carpeta, generado)
    _encabezados(hoja)
    for indice, fila in enumerate(datos):
        _fila(hoja, FILA_DE_ENCABEZADOS + 1 + indice, fila)

    ultima_fila = FILA_DE_ENCABEZADOS + len(datos)
    if datos:
        # Como tabla de Excel: franjas que se mantienen al ordenar, filtro en
        # cada encabezado y el estilo sigue si alguien añade una fila a mano.
        tabla = Table(
            displayName="InventarioDeActas",
            ref=f"A{FILA_DE_ENCABEZADOS}:{_ultima_columna()}{ultima_fila}",
        )
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
        )
        hoja.add_table(tabla)
    _nota_al_pie(hoja, ultima_fila + 2)

    # Los encabezados de la tabla quedan a la vista al bajar.
    hoja.freeze_panes = f"A{FILA_DE_ENCABEZADOS + 1}"
    _para_imprimir(hoja)

    destino.parent.mkdir(parents=True, exist_ok=True)
    libro.save(destino)
    return destino


# -- la cabecera --------------------------------------------------------------


def _banda_del_titulo(hoja, carpeta: str | None, datos: list[dict], nota: str | None) -> None:
    """El título sobre una banda marino, con una línea dorada debajo."""
    ultima = _ultima_columna()
    origen = (
        Path(carpeta).name or carpeta
        if carpeta
        else str(datos[0]["archivo"]) if len(datos) == 1 else "Actas inventariadas"
    )
    for fila, texto, fuente, alto in (
        (
            FILA_DEL_TITULO,
            "INVENTARIO DE ACTAS DE CONSEJO",
            Font(name=_FUENTE, size=18, bold=True, color="FFFFFF"),
            34,
        ),
        (
            FILA_DEL_SUBTITULO,
            "  ·  ".join(parte for parte in ("Inventario documental", origen, nota) if parte),
            Font(name=_FUENTE, size=11, color=_AZUL_CLARO),
            20,
        ),
    ):
        hoja.merge_cells(f"A{fila}:{ultima}{fila}")
        celda = hoja[f"A{fila}"]
        celda.value = texto
        celda.font = fuente
        celda.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        _rellenar(hoja, fila, 1, len(COLUMNAS), _MARINO)
        hoja.row_dimensions[fila].height = alto

    _rellenar(hoja, FILA_DE_LA_LINEA, 1, len(COLUMNAS), _DORADO)
    hoja.row_dimensions[FILA_DE_LA_LINEA].height = 4
    hoja.row_dimensions[FILA_DE_LA_LINEA + 1].height = 10


def _indicadores(hoja, datos: list[dict], carpeta: str | None, generado: datetime) -> None:
    """Una fila de tarjetas con lo que se quiere saber de un vistazo."""
    folios = sum(fila["folios"] for fila in datos if isinstance(fila.get("folios"), int))
    por_revisar = sum(1 for fila in datos if fila.get("observaciones"))
    valores = {
        "ACTAS": (len(datos), 18),
        "FOLIOS": (folios, 18),
        "PERIODO": (_periodo(datos), 12),
        "POR REVISAR": (por_revisar, 18),
        "CARPETA": (carpeta or "—", 10),
        "GENERADO": (generado.strftime("%d-%m-%Y  %H:%M"), 12),
    }
    etiqueta, valor = FILA_DE_INDICADORES, FILA_DE_INDICADORES + 1
    for nombre, desde, hasta in _INDICADORES:
        dato, tamano = valores[nombre]
        alerta = nombre == "POR REVISAR" and por_revisar > 0
        for fila in (etiqueta, valor):
            if hasta > desde:
                hoja.merge_cells(
                    start_row=fila, start_column=desde, end_row=fila, end_column=hasta
                )
            for columna in range(desde, hasta + 1):
                celda = hoja.cell(row=fila, column=columna)
                celda.fill = PatternFill("solid", fgColor=_AMBAR_FONDO if alerta else _TARJETA)
                # Un filo marino a la izquierda hace de cada indicador una
                # tarjeta; el blanco a la derecha las separa entre sí.
                celda.border = Border(
                    left=Side(style="thick", color=_AMBAR if alerta else _MARINO)
                    if columna == desde
                    else None,
                    right=Side(style="thick", color="FFFFFF") if columna == hasta else None,
                )
        arriba = hoja.cell(row=etiqueta, column=desde, value=nombre)
        arriba.font = Font(name=_FUENTE, size=8, bold=True, color=_GRIS_CLARO)
        arriba.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
        abajo = hoja.cell(row=valor, column=desde, value=dato)
        abajo.font = Font(
            name=_FUENTE, size=tamano, bold=True, color=_AMBAR if alerta else _MARINO
        )
        abajo.alignment = Alignment(
            horizontal="left", vertical="center", indent=1, shrink_to_fit=nombre == "CARPETA"
        )
    hoja.row_dimensions[etiqueta].height = 16
    hoja.row_dimensions[valor].height = 28
    hoja.row_dimensions[valor + 1].height = 12


def _periodo(datos: list[dict]) -> str:
    """De qué fecha a qué fecha van las actas, o «—» si ninguna se pudo fechar."""
    fechas = sorted(filter(None, (_fecha_de(str(fila.get("fecha") or "")) for fila in datos)))
    if not fechas:
        return "—"
    desde, hasta = fechas[0].strftime("%d-%m-%Y"), fechas[-1].strftime("%d-%m-%Y")
    return desde if desde == hasta else f"{desde}  a  {hasta}"


def _fecha_de(texto: str) -> date | None:
    """La fecha de la fila: «17-12-2012», o en letra «Lunes 17 de diciembre de 2012»."""
    cifras = _FECHA_EN_CIFRAS.fullmatch(texto.strip())
    if cifras:
        dia, mes, anio = (int(parte) for parte in cifras.groups())
        try:
            return date(anio, mes, dia)
        except ValueError:
            return None
    encontrada = _FECHA_LARGA.search(texto)
    if not encontrada:
        return None
    dia, mes, anio = encontrada.groups()
    numero = MESES.get(mes.lower().translate(str.maketrans("áéíóú", "aeiou")))
    if numero is None:
        return None
    try:
        return date(int(anio), numero, int(dia))
    except ValueError:
        return None


# -- la tabla -----------------------------------------------------------------


def _encabezados(hoja) -> None:
    for columna, (titulo, _clave, _ancho) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=FILA_DE_ENCABEZADOS, column=columna, value=titulo)
        celda.font = Font(name=_FUENTE, size=10, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=_MARINO)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = Border(
            left=Side(style="thin", color=_MARINO),
            right=Side(style="thin", color=_MARINO),
            top=Side(style="thin", color=_MARINO),
            bottom=Side(style="medium", color=_DORADO),
        )
    hoja.row_dimensions[FILA_DE_ENCABEZADOS].height = 32


def _fila(hoja, numero: int, datos: dict) -> None:
    renglones = 1
    for columna, (_titulo, clave, ancho) in enumerate(COLUMNAS, start=1):
        valor = datos.get(clave, "")
        celda = hoja.cell(row=numero, column=columna, value=valor)
        centrada = clave in _CENTRADAS
        celda.border = _CAJA
        celda.alignment = Alignment(
            horizontal="center" if centrada else "left",
            vertical="center",
            wrap_text=not centrada,
        )
        # Lo que no se pudo leer, en gris: está, pero no compite con lo leído.
        sin_dato = valor in ("N/A", "", None)
        if clave == "observaciones" and valor:
            # Lo que hay que revisar se distingue sin tener que leer la fila.
            celda.font = Font(name=_FUENTE, size=9, italic=True, color=_AMBAR)
            celda.fill = PatternFill("solid", fgColor=_AMBAR_FONDO)
        elif clave == "acta" and not sin_dato:
            celda.font = Font(name=_FUENTE, size=10, bold=True, color=_MARINO)
        else:
            celda.font = Font(name=_FUENTE, size=10, color=_GRIS_CLARO if sin_dato else None)
        if not centrada and valor:
            renglones = max(renglones, _renglones(str(valor), ancho))
    # Excel no recalcula el alto de un renglón con texto ajustado que no tocó
    # nadie: sin esto, el resumen se vería cortado en una sola línea.
    hoja.row_dimensions[numero].height = min(409, max(20, renglones * _ALTO_DE_RENGLON + 6))


def _nota_al_pie(hoja, fila: int) -> None:
    """De dónde sale cada dato, dicho una vez debajo de la tabla."""
    hoja.merge_cells(f"A{fila}:{_ultima_columna()}{fila}")
    celda = hoja[f"A{fila}"]
    celda.value = (
        "El acta y la fecha se leen del encabezado impreso del acta. El folio inicial y el "
        "final son los números escritos a mano en la esquina superior derecha de la primera y "
        "la última hoja foliada, y el número de folios cuenta las hojas entre los dos. El "
        "resumen abarca el documento completo. N/A indica un dato que no se pudo leer; la "
        "columna Observaciones explica por qué."
    )
    celda.font = Font(name=_FUENTE, size=9, italic=True, color=_GRIS)
    celda.alignment = Alignment(wrap_text=True, vertical="top")
    hoja.row_dimensions[fila].height = 30


# -- la impresión -------------------------------------------------------------


def _para_imprimir(hoja) -> None:
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_LETTER
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_options.horizontalCentered = True
    hoja.print_title_rows = f"{FILA_DE_ENCABEZADOS}:{FILA_DE_ENCABEZADOS}"
    hoja.page_margins.left = hoja.page_margins.right = 0.4
    hoja.page_margins.top = hoja.page_margins.bottom = 0.55
    hoja.oddFooter.left.text = PROGRAMA
    hoja.oddFooter.left.size = 8
    hoja.oddFooter.right.text = "Página &P de &N"
    hoja.oddFooter.right.size = 8


# -- utilidades ---------------------------------------------------------------


def _ultima_columna() -> str:
    return get_column_letter(len(COLUMNAS))


def _rellenar(hoja, fila: int, desde: int, hasta: int, color: str) -> None:
    for columna in range(desde, hasta + 1):
        hoja.cell(row=fila, column=columna).fill = PatternFill("solid", fgColor=color)


def _renglones(texto: str, ancho: int) -> int:
    """Cuántos renglones ocupa el texto en una columna de ese ancho."""
    caracteres = max(1, int(ancho * 1.15))
    return sum(max(1, math.ceil(len(parrafo) / caracteres)) for parrafo in texto.splitlines() or [""])


def _paginas(fila: Mapping[str, object]) -> str:
    """«5–29»: qué páginas del PDF son esta acta. En un acta suelta, todas."""
    desde, hasta = fila.get("pagina_inicial"), fila.get("pagina_final")
    if not isinstance(desde, int) or not isinstance(hasta, int):
        return "—"
    return str(desde) if desde == hasta else f"{desde}–{hasta}"


def _como_dict(fila: Mapping[str, object] | object) -> dict[str, object]:
    if isinstance(fila, Mapping):
        return dict(fila)
    if is_dataclass(fila) and not isinstance(fila, type):
        return asdict(fila)
    raise TypeError(f"no sé leer una fila de tipo {type(fila).__name__}")
