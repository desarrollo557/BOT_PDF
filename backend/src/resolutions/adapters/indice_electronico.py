"""El Índice Electrónico del expediente, escrito sobre la plantilla del cliente.

La plantilla es la que entregó el operador -- «00.INDICE ELECTRONICO.xlsx» --,
copiada tal cual en ``assets/plantillas``: título, «Serie o Subserie
Documental: HISTORIALES DE CLIENTES», los rótulos, la letra, los bordes y el
pie «FECHA DE CIERRE DEL EXPEDIENTE:». Aquí no se dibuja nada: se escriben los
datos en sus celdas y se copia el formato de su primera fila de documentos a
las que hagan falta, para que el archivo salga igual que el que se hace a mano.

Las fechas van como texto «dd/mm/aaaa», que es como las lleva la plantilla.
"""

from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

#: La plantilla del cliente.
PLANTILLA = Path(__file__).resolve().parents[1] / "assets" / "plantillas" / "indice_electronico.xlsx"
#: Con qué nombre se deja junto a los PDF: el mismo de la plantilla.
NOMBRE_DEL_INDICE = "00.INDICE ELECTRONICO.xlsx"
#: Dónde está cada cosa en la plantilla.
CELDA_NOMBRE_DEL_EXPEDIENTE = "B3"
CELDA_NIC = "B4"
CELDA_TERRITORIAL = "B5"
PRIMERA_FILA = 8
COLUMNAS = "ABCDEFGH"
#: Lo que dice la celda que cierra la tabla.
CIERRE = "FECHA DE CIERRE DEL EXPEDIENTE"
#: Lo que se escribe donde no hay dato, como hace la plantilla con el nombre
#: del expediente: una raya dice «no aplica o no se supo», una celda vacía no
#: dice nada.
SIN_DATO = "-"


@dataclass(frozen=True, slots=True)
class FilaDelIndice:
    """Un documento del expediente, con lo que pide cada columna."""

    nombre: str
    #: La fecha que se ve en el documento -- su sello, su guía --; ``None``
    #: si no trae ninguna.
    fecha_creacion: date | None
    #: El día en que el bot lo incorporó al expediente.
    fecha_incorporacion: date | None
    orden: int
    paginas: int
    formato: str = "pdf"
    observaciones: str = ""
    responsable: str = ""


#: Las columnas que crecen con su texto, y cuánto.
COLUMNA_DEL_NOMBRE = "A"
COLUMNA_DE_OBSERVACIONES = "G"
ANCHO_DE_OBSERVACIONES = 80
#: Alto de un renglón de letra 12, en puntos.
ALTO_DE_RENGLON = 15.5


def _renglones(texto: str, ancho: float) -> int:
    """Cuántos renglones ocupa un texto ajustado a una columna de ese ancho."""
    if not texto:
        return 1
    por_renglon = max(int(ancho * 1.05), 10)
    return sum(max(1, -(-len(parte) // por_renglon)) for parte in str(texto).split("\n"))


def _fecha(valor: date | None) -> str:
    return valor.strftime("%d/%m/%Y") if valor else "S.F."


def escribir_indice(
    destino: Path,
    filas: list[FilaDelIndice],
    *,
    nic: str | None = None,
    territorial: str | None,
    nombre_del_expediente: str | None = None,
    plantilla: Path = PLANTILLA,
) -> Path:
    """Escribe el índice en ``destino`` y devuelve la ruta del archivo."""
    from openpyxl import load_workbook

    libro = load_workbook(plantilla)
    _llenar(
        libro.active,
        filas,
        nic=nic,
        territorial=territorial,
        nombre_del_expediente=nombre_del_expediente,
    )
    destino.mkdir(parents=True, exist_ok=True)
    ruta = destino / NOMBRE_DEL_INDICE
    libro.save(ruta)
    return ruta


@dataclass(frozen=True, slots=True)
class IndiceLeido:
    """Un índice ya entregado, leído de vuelta: su encabezado y sus filas."""

    filas: list[FilaDelIndice]
    nic: str | None
    territorial: str | None
    nombre_del_expediente: str | None


def _fecha_leida(valor: object) -> date | None:
    try:
        return datetime.strptime(str(valor).strip(), "%d/%m/%Y").date()
    except ValueError:
        return None


def _dato(valor: object) -> str | None:
    texto = str(valor).strip() if valor is not None else ""
    return None if texto in ("", SIN_DATO) else texto


def leer_indice(ruta: Path) -> IndiceLeido:
    """Lo que dice un índice escrito por :func:`escribir_indice`.

    Sirve para juntar los de varias cajas sin volver a calcular nada: lo que
    sale en el libro del lote es exactamente lo que se entregó con cada caja.
    """
    from openpyxl import load_workbook

    # Cerrado al terminar: en modo de sólo lectura el libro deja el archivo
    # abierto, y en Windows un archivo abierto no se puede mover ni apartar.
    libro = load_workbook(ruta, read_only=True)
    try:
        hoja = libro.active
        filas: list[FilaDelIndice] = []
        fila = PRIMERA_FILA
        while True:
            valores = [hoja[f"{c}{fila}"].value for c in COLUMNAS]
            nombre = str(valores[0] or "").strip()
            if not nombre or nombre.upper().startswith(CIERRE):
                break
            filas.append(
                FilaDelIndice(
                    nombre=nombre,
                    fecha_creacion=_fecha_leida(valores[1]),
                    fecha_incorporacion=_fecha_leida(valores[2]),
                    orden=int(valores[3] or len(filas) + 1),
                    paginas=int(valores[4] or 0),
                    formato=str(valores[5] or "pdf"),
                    observaciones=str(valores[6] or ""),
                    responsable=str(valores[7] or ""),
                )
            )
            fila += 1
        return IndiceLeido(
            filas=filas,
            nic=_dato(hoja[CELDA_NIC].value),
            territorial=_dato(hoja[CELDA_TERRITORIAL].value),
            nombre_del_expediente=_dato(hoja[CELDA_NOMBRE_DEL_EXPEDIENTE].value),
        )
    finally:
        libro.close()


#: Lo más largo que Excel admite en el nombre de una hoja.
LARGO_DE_HOJA = 31


def escribir_indices_por_hoja(
    destino: Path,
    cajas: list[tuple[str, IndiceLeido]],
    *,
    nombre_del_archivo: str = NOMBRE_DEL_INDICE,
    plantilla: Path = PLANTILLA,
) -> Path:
    """Un libro con una hoja por caja, cada una con la plantilla del cliente.

    Es el inventario general de un lote: lo pidió el operador -- «una hoja
    diferente por PDF procesado» --, así que cada hoja es el Índice Electrónico
    de una caja tal como se entregó junto a sus PDF, con su propio orden.
    """
    from openpyxl import load_workbook

    libro = load_workbook(plantilla)
    base = libro.active
    usados: set[str] = set()
    for nombre, indice in cajas:
        hoja = libro.copy_worksheet(base)
        hoja.title = _nombre_de_hoja(nombre, usados)
        hoja.page_setup.orientation = base.page_setup.orientation
        hoja.sheet_properties.pageSetUpPr = copy(base.sheet_properties.pageSetUpPr)
        hoja.page_setup.fitToWidth = base.page_setup.fitToWidth
        hoja.page_setup.fitToHeight = base.page_setup.fitToHeight
        _llenar(
            hoja,
            indice.filas,
            nic=indice.nic,
            territorial=indice.territorial,
            nombre_del_expediente=indice.nombre_del_expediente or nombre,
        )
    if cajas:
        libro.remove(base)
        libro.active = 0
    destino.mkdir(parents=True, exist_ok=True)
    ruta = destino / nombre_del_archivo
    libro.save(ruta)
    return ruta


def _nombre_de_hoja(nombre: str, usados: set[str]) -> str:
    """El nombre de la caja como nombre de hoja: sin lo que Excel no admite, y sin repetir."""
    limpio = re.sub(r"[\[\]:*?/\\]+", " ", nombre).strip() or "Caja"
    limpio = limpio[:LARGO_DE_HOJA]
    candidato, n = limpio, 2
    while candidato.lower() in usados:
        sufijo = f" ({n})"
        candidato = limpio[: LARGO_DE_HOJA - len(sufijo)] + sufijo
        n += 1
    usados.add(candidato.lower())
    return candidato


def _llenar(
    hoja,
    filas: list[FilaDelIndice],
    *,
    nic: str | None,
    territorial: str | None,
    nombre_del_expediente: str | None,
) -> None:
    """Escribe el encabezado y las filas en una hoja que es copia de la plantilla."""
    from openpyxl.styles import Alignment

    hoja[CELDA_NOMBRE_DEL_EXPEDIENTE] = nombre_del_expediente or SIN_DATO
    # Sin NIC la celda queda vacía, no con raya: se llena a mano.
    hoja[CELDA_NIC] = int(nic) if nic and nic.isdigit() else (nic or None)
    hoja[CELDA_TERRITORIAL] = territorial or SIN_DATO

    # El formato de una fila de documento y el del pie, tomados de la
    # plantilla antes de tocarla.
    estilo = [copy(hoja[f"{c}{PRIMERA_FILA}"]._style) for c in COLUMNAS]
    alto = hoja.row_dimensions[PRIMERA_FILA].height
    fila_del_cierre = next(
        fila
        for fila in range(PRIMERA_FILA, hoja.max_row + 1)
        if str(hoja[f"A{fila}"].value or "").strip().upper().startswith(CIERRE)
    )
    estilo_del_cierre = [copy(hoja[f"{c}{fila_del_cierre}"]._style) for c in COLUMNAS]
    alto_del_cierre = hoja.row_dimensions[fila_del_cierre].height
    texto_del_cierre = hoja[f"A{fila_del_cierre}"].value

    # Fuera las filas de ejemplo y el pie, que se vuelve a poner debajo.
    for rango in list(hoja.merged_cells.ranges):
        if rango.min_row >= PRIMERA_FILA:
            hoja.unmerge_cells(str(rango))
    for fila in range(PRIMERA_FILA, hoja.max_row + 1):
        for c in COLUMNAS:
            hoja[f"{c}{fila}"].value = None

    for desplazamiento, documento in enumerate(filas):
        fila = PRIMERA_FILA + desplazamiento
        valores = (
            documento.nombre,
            _fecha(documento.fecha_creacion),
            _fecha(documento.fecha_incorporacion),
            documento.orden,
            documento.paginas,
            documento.formato,
            documento.observaciones or None,
            documento.responsable or None,
        )
        for c, valor, formato in zip(COLUMNAS, valores, estilo, strict=True):
            celda = hoja[f"{c}{fila}"]
            celda.value = valor
            celda._style = copy(formato)
        # Las observaciones traen el detalle del trámite -- qué documentos
        # junta, sus fechas, sus anexos -- y en la columna de la plantilla, de
        # trece caracteres, no se leía nada. Se ajusta el texto y la fila crece
        # lo que haga falta; el nombre también, que lleva el número RE.
        for c in (COLUMNA_DEL_NOMBRE, COLUMNA_DE_OBSERVACIONES):
            celda = hoja[f"{c}{fila}"]
            celda.alignment = Alignment(
                horizontal="left" if c == COLUMNA_DE_OBSERVACIONES else celda.alignment.horizontal,
                vertical="center",
                wrap_text=True,
            )
        renglones = max(
            _renglones(documento.observaciones, ANCHO_DE_OBSERVACIONES),
            _renglones(documento.nombre, hoja.column_dimensions[COLUMNA_DEL_NOMBRE].width or 70),
        )
        hoja.row_dimensions[fila].height = max(alto or 0, ALTO_DE_RENGLON * renglones + 6)
    if any(documento.observaciones for documento in filas):
        hoja.column_dimensions[COLUMNA_DE_OBSERVACIONES].width = max(
            hoja.column_dimensions[COLUMNA_DE_OBSERVACIONES].width or 0, ANCHO_DE_OBSERVACIONES
        )

    cierre = PRIMERA_FILA + len(filas)
    for c, formato in zip(COLUMNAS, estilo_del_cierre, strict=True):
        hoja[f"{c}{cierre}"]._style = copy(formato)
    hoja[f"A{cierre}"] = texto_del_cierre
    hoja.merge_cells(f"A{cierre}:H{cierre}")
    hoja.row_dimensions[cierre].height = alto_del_cierre
    hoja.print_area = f"A1:H{cierre}"
