"""Escribe el Formato Único de Inventario Documental (FUID) de la Universidad.

No genera una planilla propia: abre la plantilla oficial -- FO-GD-008, versión
00 -- y la rellena. El logotipo, la banda de la cabecera, los encabezados a dos
niveles y el bloque de firmas son los del formato aprobado, y lo único que este
módulo aporta son las filas y el sitio al que se corre la firma cuando las filas
no caben en el hueco que la plantilla deja.

Las reglas de contenido no son de este autor: salen del instructivo del propio
formato, que dice qué va en cada columna y, lo más importante, que **lo que no
aplica se escribe N/A**. Una celda vacía en un inventario documental no dice
"no aplica", dice "se olvidó".
"""

from __future__ import annotations

import logging
from copy import copy
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..application.fuid import Cabecera, FuidRow
from ..application.progress import (
    NullProgressReporter,
    ProgressEvent,
    ProgressReporter,
    Stage,
)
from .guardar_libro import guardar_libro

logger = logging.getLogger(__name__)

#: La columna NOTAS en la fila de ``FuidRow.as_cells``, la última.
COLUMNA_DE_NOTAS = 15


HOJA = "INVENTARIO"

#: La plantilla congela los paneles en A14, que es como declara dónde empiezan
#: los datos. La fila 14 que trae escrita es la guía de qué va en cada columna,
#: y el primer registro la sustituye.
PRIMERA_FILA = 14

#: De dónde se copia el formato de cada fila de datos. Es la misma fila 14: el
#: estilo se lee antes de escribir nada encima.
FILA_ESTILO = 14

#: Filas de la cabecera que identifican la entrega.
FILA_OFICINA = 8
FILA_OBJETO = 9
COLUMNA_CABECERA = 2

#: Dónde está el bloque de firmas en la plantilla en blanco. Es sólo el punto
#: de partida: se busca por su texto, porque una plantilla que ya se usó lo
#: tiene mucho más abajo -- la de diplomas, con 1.084 filas dentro, lo lleva en
#: la 1100 -- y buscarlo en una fila fija encontraría datos y los tomaría por
#: firmas.
BLOQUE_FIRMAS = (109, 112)

#: Con qué empieza la primera línea del bloque. Es lo que lo identifica.
MARCA_DE_FIRMAS = "elaborado por"

#: Cuántas líneas ocupa: elaborado / cargo / firma / lugar.
ALTO_DE_FIRMAS = 4

FILAS_ANTES_DE_FIRMAR = 2

#: Columnas A a P. Las dos siguientes que trae la hoja están fuera del formato.
COLUMNAS = 16

#: Las columnas que van a la derecha de NOTAS (Q a T) cuando alguna fila es un
#: acta de consejo: su título y su ancho. Las pidió el operador para el
#: inventario de Odontología el 2026-10-31. Las dos últimas van bajo un mismo
#: título, como «CONSECUTIVO» o «FECHAS EXTREMAS» en el formato.
COLUMNAS_DE_ACTAS = (("ACTA #", 10), ("OBJETIVO / RESUMEN", 70), ("Desde", 9), ("Hasta", 9))
TITULO_DE_LAS_PAGINAS = "PÁGINAS DEL PDF"
FILA_DE_TITULOS = 12
FILA_DE_SUBTITULOS = 13
#: De qué celdas de la plantilla se copia el estilo: el título de NOTAS, el
#: subtítulo «Inicial» de CONSECUTIVO, y en las filas, NOTAS y No. de Folios.
_ESTILO_DE_TITULO = "P12"
_ESTILO_DE_SUBTITULO = "D13"
_COLUMNA_DE_FOLIOS = 11
#: Cuántos caracteres caben en un renglón de la columna del resumen, y el alto
#: de cada renglón: la fila crece para que el resumen se lea entero.
_LETRAS_POR_RENGLON = 80
_ALTO_DEL_RENGLON = 13

#: Cada cuántas filas se dice por dónde va. Escribir un FUID de 1.084 filas
#: tarda cinco segundos y medio, y son los últimos del trabajo: la barra de
#: páginas ya está al 100 % y sin este aviso la pantalla parece haberse quedado
#: colgada justo al final. Cincuenta filas es un aviso cada cuarto de segundo,
#: que se nota sin inundar nada.
FILAS_POR_AVISO = 50


def _combinar(hoja: Worksheet, fila: int, columna: int, hasta_fila: int, hasta_columna: int) -> None:
    """Combina las celdas, salvo que ya lo estén: la plantilla puede ser un FUID ya escrito."""
    from openpyxl.utils import get_column_letter

    rango = f"{get_column_letter(columna)}{fila}:{get_column_letter(hasta_columna)}{hasta_fila}"
    if rango not in {str(combinado) for combinado in hoja.merged_cells.ranges}:
        hoja.merge_cells(rango)


@dataclass
class _Firma:
    """Una celda del bloque de firmas, guardada para volver a colocarla."""

    fila_relativa: int
    columna: int
    valor: object
    estilo: object
    ancho_combinado: int = 1


class FuidInventory:
    """Rellena la plantilla oficial con las filas que se le den."""

    def __init__(
        self, plantilla: Path, progress: ProgressReporter | None = None
    ) -> None:
        self._plantilla = Path(plantilla)
        if not self._plantilla.is_file():
            raise FileNotFoundError(f"no se encuentra la plantilla del FUID: {plantilla}")
        self._progress = progress or NullProgressReporter()

    def write(
        self,
        filas: list[FuidRow],
        destino: Path,
        *,
        cabecera: Cabecera | None = None,
    ) -> Path:
        total = len(filas)
        # Abrir la plantilla es lo que más tarda de todo esto -- openpyxl lee el
        # libro entero, con sus estilos -- y ocurre antes de la primera fila, así
        # que hay que decirlo antes de empezar o parecerá que no pasa nada.
        self._anunciar(0, total, "abriendo la plantilla del formato")
        libro = load_workbook(self._plantilla)
        hoja = libro[HOJA]

        self._cabecera(hoja, cabecera or Cabecera())
        estilos = self._estilos_de_fila(hoja)
        desde_firmas, firmas = self._recoger_firmas(hoja)
        con_actas = any(fila.acta is not None for fila in filas)
        if con_actas:
            self._cabecera_de_actas(hoja)

        ultima = PRIMERA_FILA + len(filas) - 1
        for indice, fila in enumerate(filas):
            self._escribir_fila(hoja, PRIMERA_FILA + indice, fila, estilos, con_actas=con_actas)
            if (indice + 1) % FILAS_POR_AVISO == 0:
                self._anunciar(indice + 1, total, "escribiendo las filas")

        # Todo lo que quede por debajo del último registro se limpia. No es sólo
        # cosmética: cuando la plantilla es un inventario anterior -- que es
        # cómo se usa la de diplomas -- ahí abajo hay cientos de filas de otra
        # entrega, y dejarlas convertiría el documento en dos inventarios
        # pegados sin que nadie lo note.
        for numero in range(ultima + 1, max(desde_firmas, ultima) + 1):
            self._vaciar(hoja, numero)

        self._colocar_firmas(hoja, firmas, ultima + FILAS_ANTES_DE_FIRMAR + 1)

        # Guardar vuelve a recorrer el libro entero. Es el otro tramo largo, y
        # el último: si no se anuncia, el trabajo parece colgado en la fila final.
        self._anunciar(total, total, "guardando la planilla")
        destino.parent.mkdir(parents=True, exist_ok=True)
        guardar_libro(libro, destino)
        return destino

    def _anunciar(self, hechas: int, total: int, detalle: str) -> None:
        """Decir por dónde va. Un fallo del aviso nunca pierde la planilla."""
        try:
            self._progress.emit(
                ProgressEvent(
                    stage=Stage.INVENTORYING, done=hechas, total=total, detail=detalle
                )
            )
        except Exception:  # noqa: BLE001 - la telemetría nunca rompe el trabajo
            # Se traga, pero deja rastro: un aviso que falla en silencio hace
            # que la barra se quede quieta sin que nadie sepa por qué.
            logger.debug("no se pudo anunciar el avance del FUID", exc_info=True)

    # -- cabecera -------------------------------------------------------------

    @staticmethod
    def _cabecera(hoja: Worksheet, cabecera: Cabecera) -> None:
        if cabecera.oficina_productora:
            hoja.cell(row=FILA_OFICINA, column=COLUMNA_CABECERA).value = (
                f"OFICINA PRODUCTORA: {cabecera.oficina_productora}"
            )
        if cabecera.objeto:
            hoja.cell(row=FILA_OBJETO, column=COLUMNA_CABECERA).value = (
                f"OBJETO: {cabecera.objeto}"
            )

    @staticmethod
    def _cabecera_de_actas(hoja: Worksheet) -> None:
        """Los títulos de las columnas de actas, con el estilo de los del formato."""
        titulo, subtitulo = hoja[_ESTILO_DE_TITULO], hoja[_ESTILO_DE_SUBTITULO]
        primera = COLUMNAS + 1
        for desplazamiento, (nombre, ancho) in enumerate(COLUMNAS_DE_ACTAS):
            columna = primera + desplazamiento
            letra = hoja.cell(row=FILA_DE_TITULOS, column=columna).column_letter
            hoja.column_dimensions[letra].width = ancho
            de_paginas = desplazamiento >= 2
            arriba = hoja.cell(row=FILA_DE_TITULOS, column=columna)
            abajo = hoja.cell(row=FILA_DE_SUBTITULOS, column=columna)
            for celda, modelo in ((arriba, titulo), (abajo, subtitulo if de_paginas else titulo)):
                celda.font = copy(modelo.font)
                celda.border = copy(modelo.border)
                celda.fill = copy(modelo.fill)
                celda.alignment = copy(modelo.alignment)
            if de_paginas:
                abajo.value = nombre
            else:
                arriba.value = nombre
                _combinar(hoja, FILA_DE_TITULOS, columna, FILA_DE_SUBTITULOS, columna)
        hoja.cell(row=FILA_DE_TITULOS, column=primera + 2).value = TITULO_DE_LAS_PAGINAS
        _combinar(hoja, FILA_DE_TITULOS, primera + 2, FILA_DE_TITULOS, primera + 3)

    # -- filas ----------------------------------------------------------------

    @staticmethod
    def _estilos_de_fila(hoja: Worksheet) -> list[dict]:
        """El formato de la fila 14, que es el que llevan todos los registros.

        Se copia celda a celda en lugar de aplicar uno inventado, para que la
        planilla generada sea indistinguible de una rellenada a mano sobre la
        misma plantilla.
        """
        estilos = []
        for columna in range(1, COLUMNAS + 1):
            celda = hoja.cell(row=FILA_ESTILO, column=columna)
            estilos.append(
                {
                    "font": copy(celda.font),
                    "border": copy(celda.border),
                    "fill": copy(celda.fill),
                    "alignment": copy(celda.alignment),
                    "number_format": celda.number_format,
                }
            )
        return estilos

    @staticmethod
    def _escribir_fila(
        hoja: Worksheet, numero: int, fila: FuidRow, estilos: list[dict], *, con_actas: bool = False
    ) -> None:
        for indice, valor in enumerate(fila.as_cells()):
            if indice == COLUMNA_DE_NOTAS:
                # NOTAS va vacía en el FUID que se entrega al cliente: lo pidió
                # el operador el 2026-10-30 para todos los del FO-GD-008. Lo
                # que no se pudo leer sigue en la cola de revisión del bot.
                valor = None
            celda = hoja.cell(row=numero, column=indice + 1, value=valor)
            estilo = estilos[indice]
            celda.font = copy(estilo["font"])
            celda.border = copy(estilo["border"])
            celda.fill = copy(estilo["fill"])
            celda.alignment = copy(estilo["alignment"])
            # El número de orden y los consecutivos son texto en la plantilla:
            # un folio "336BIS" y una cédula con ceros a la izquierda dejan de
            # ser lo que son en cuanto Excel los convierte en número.
            celda.number_format = estilo["number_format"]
        hoja.row_dimensions[numero].height = 20
        if con_actas:
            FuidInventory._escribir_acta(hoja, numero, fila, estilos)

    @staticmethod
    def _escribir_acta(hoja: Worksheet, numero: int, fila: FuidRow, estilos: list[dict]) -> None:
        """ACTA #, OBJETIVO / RESUMEN y las páginas desde y hasta, a la derecha de NOTAS."""
        from openpyxl.styles import Alignment

        for desplazamiento, valor in enumerate(fila.celdas_del_acta()):
            celda = hoja.cell(row=numero, column=COLUMNAS + 1 + desplazamiento, value=valor)
            es_resumen = desplazamiento == 1
            estilo = estilos[COLUMNA_DE_NOTAS if es_resumen else _COLUMNA_DE_FOLIOS]
            celda.font = copy(estilo["font"])
            celda.border = copy(estilo["border"])
            celda.fill = copy(estilo["fill"])
            if es_resumen:
                celda.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                celda.alignment = Alignment(horizontal="center", vertical="top")
        resumen = fila.resumen or ""
        renglones = -(-len(resumen) // _LETRAS_POR_RENGLON)
        if renglones > 1:
            hoja.row_dimensions[numero].height = max(20, renglones * _ALTO_DEL_RENGLON)

    @staticmethod
    def _vaciar(hoja: Worksheet, numero: int) -> None:
        """Deja la fila en blanco y sin cuadrícula.

        La fila de la que se copia el estilo es la excepción: pierde el valor
        pero conserva el formato. Es lo que permite que una planilla ya escrita
        vuelva a servir de plantilla -- sin esto, guardar un inventario vacío se
        llevaría por delante el formato de todas las filas siguientes.
        """
        from openpyxl.styles import Border

        conserva_el_formato = numero == FILA_ESTILO
        for columna in range(1, COLUMNAS + len(COLUMNAS_DE_ACTAS) + 1):
            celda = hoja.cell(row=numero, column=columna)
            celda.value = None
            if not conserva_el_formato:
                celda.border = Border()

    # -- firmas ---------------------------------------------------------------

    @staticmethod
    def _localizar_firmas(hoja: Worksheet) -> tuple[int, int]:
        """Encuentra el bloque de firmas por lo que dice, no por dónde está.

        Una plantilla en blanco lo tiene justo debajo del área de datos; una que
        ya se rellenó, cientos de filas más abajo. Buscarlo por su texto es lo
        único que funciona en las dos.
        """
        for numero in range(PRIMERA_FILA, hoja.max_row + 1):
            valor = hoja.cell(row=numero, column=1).value
            if isinstance(valor, str) and valor.strip().lower().startswith(MARCA_DE_FIRMAS):
                return numero, numero + ALTO_DE_FIRMAS - 1
        return BLOQUE_FIRMAS

    @staticmethod
    def _recoger_firmas(hoja: Worksheet) -> tuple[int, list[_Firma]]:
        """Guarda el bloque de firmas y lo borra de donde estaba."""
        inicio, fin = FuidInventory._localizar_firmas(hoja)
        anchos: dict[tuple[int, int], int] = {}
        for rango in list(hoja.merged_cells.ranges):
            if inicio <= rango.min_row <= fin:
                anchos[(rango.min_row, rango.min_col)] = rango.max_col - rango.min_col + 1
                hoja.unmerge_cells(str(rango))

        recogidas: list[_Firma] = []
        for numero in range(inicio, fin + 1):
            for columna in range(1, COLUMNAS + 1):
                celda = hoja.cell(row=numero, column=columna)
                if celda.value in (None, ""):
                    continue
                recogidas.append(
                    _Firma(
                        fila_relativa=numero - inicio,
                        columna=columna,
                        valor=celda.value,
                        estilo=copy(celda._style),
                        ancho_combinado=anchos.get((numero, columna), 1),
                    )
                )
                celda.value = None

        return fin, recogidas

    @staticmethod
    def _colocar_firmas(hoja: Worksheet, firmas: list[_Firma], desde: int) -> None:
        for firma in firmas:
            numero = desde + firma.fila_relativa
            celda = hoja.cell(row=numero, column=firma.columna, value=firma.valor)
            celda._style = copy(firma.estilo)
            if firma.ancho_combinado > 1:
                hoja.merge_cells(
                    start_row=numero,
                    start_column=firma.columna,
                    end_row=numero,
                    end_column=firma.columna + firma.ancho_combinado - 1,
                )
