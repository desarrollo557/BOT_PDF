"""El análisis de la entrega de una carpeta: qué se procesó, qué no y por qué.

Lo pidió el operador el 2026-10-22: «estructúralo bien para que sea fácil
reconocer las carpetas y los PDF ya procesados y cuáles no… que me saque un
análisis de los procesados, los no procesados». Va en la raíz del destino,
«00.ANALISIS DE LA ENTREGA.xlsx», con cuatro hojas:

- **Resumen**: cuántos PDF, cuántos ya estaban entregados, cuántos se
  procesaron, se rehicieron, fallaron o quedan pendientes, y cuántas carpetas
  quedaron completas.
- **Por carpeta**: cada carpeta (el NIC), con sus cuentas y su estado.
- **Por PDF**: cada PDF, con lo que se decidió al verificarlo, por qué, y lo
  que pasó después.
- **Incidencias**: carpetas que no se pudieron leer y otros avisos que impiden
  interpretar la corrida como completa.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from .guardar_libro import guardar_libro

NOMBRE_DEL_ANALISIS = "00.ANALISIS DE LA ENTREGA.xlsx"

#: Lo que le pasó a cada PDF en la corrida.
YA_ENTREGADO = "Ya estaba entregado"
PROCESADO = "Procesado"
FALLO = "Falló"
EN_PROCESO = "En proceso"
PENDIENTE = "Pendiente"
REVISADO = "Sólo revisado"
CANCELADO = "Cancelado"


@dataclass(frozen=True, slots=True)
class FilaDelAnalisis:
    carpeta: str
    pdf: str
    verificacion: str
    motivo: str
    resultado: str
    detalle: str = ""
    documentos: int = 0
    paginas: int = 0


def estado_de_la_carpeta(filas: list[FilaDelAnalisis]) -> str:
    """Completa, con fallos o con pendientes: lo primero que hay que mirar."""
    resultados = {fila.resultado for fila in filas}
    if FALLO in resultados:
        return "Con fallos"
    if resultados & {PENDIENTE, EN_PROCESO, CANCELADO}:
        return "Con pendientes"
    return "Completa"


def escribir_analisis(
    destino: Path,
    filas: list[FilaDelAnalisis],
    *,
    origen: str,
    fecha: str,
    avisos: list[str] | None = None,
) -> Path:
    """Escribe el análisis en ``destino`` (entero o nada) y devuelve su ruta."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    negrita = Font(bold=True)
    cabecera = PatternFill("solid", fgColor="DDE6F0")
    colores = {
        YA_ENTREGADO: "E2EFDA",
        PROCESADO: "C6EFCE",
        FALLO: "F8CBAD",
        PENDIENTE: "FFF2CC",
        EN_PROCESO: "FFF2CC",
        CANCELADO: "FFF2CC",
        "Completa": "C6EFCE",
        "Con fallos": "F8CBAD",
        "Con pendientes": "FFF2CC",
    }

    def tabla(hoja, columnas: list[tuple[str, int]], renglones: list[list], color_en: int | None = None) -> None:
        hoja.append([nombre for nombre, _ in columnas])
        for celda in hoja[1]:
            celda.font, celda.fill = negrita, cabecera
        for indice, (_, ancho) in enumerate(columnas, start=1):
            hoja.column_dimensions[hoja.cell(1, indice).column_letter].width = ancho
        for renglon in renglones:
            hoja.append(renglon)
            if color_en is not None and (color := colores.get(str(renglon[color_en]))):
                hoja.cell(hoja.max_row, color_en + 1).fill = PatternFill("solid", fgColor=color)
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions

    filas = sorted(filas, key=lambda f: (f.carpeta.lower(), f.pdf.lower()))
    avisos = avisos or []
    por_carpeta: dict[str, list[FilaDelAnalisis]] = {}
    for fila in filas:
        por_carpeta.setdefault(fila.carpeta, []).append(fila)
    resultados = Counter(fila.resultado for fila in filas)
    decisiones = Counter(fila.verificacion for fila in filas)
    estados = Counter(estado_de_la_carpeta(suyas) for suyas in por_carpeta.values())

    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    renglones = [
        ["Origen", origen],
        ["Fecha del análisis", fecha],
        ["", ""],
        ["PDF vistos", len(filas)],
        *[
            [f"PDF · {nombre}", resultados.get(nombre, 0)]
            for nombre in (YA_ENTREGADO, PROCESADO, REVISADO, FALLO, CANCELADO, EN_PROCESO, PENDIENTE)
        ],
        ["", ""],
        *[[f"Verificación · {nombre}", cuantos] for nombre, cuantos in sorted(decisiones.items())],
        ["", ""],
        ["Incidencias", len(avisos)],
        ["", ""],
        ["Carpetas", len(por_carpeta)],
        *[[f"Carpetas · {nombre}", estados.get(nombre, 0)] for nombre in ("Completa", "Con pendientes", "Con fallos")],
    ]
    tabla(resumen, [("Concepto", 42), ("Cantidad", 60)], renglones)

    tabla(
        libro.create_sheet("Por carpeta"),
        [("Carpeta", 22), ("Estado", 16), ("PDF", 8), ("Ya entregados", 14), ("Procesados", 12),
         ("Fallidos", 10), ("Pendientes o cancelados", 22), ("Documentos entregados", 22)],
        [
            [
                carpeta,
                estado_de_la_carpeta(suyas),
                len(suyas),
                sum(f.resultado == YA_ENTREGADO for f in suyas),
                sum(f.resultado == PROCESADO for f in suyas),
                sum(f.resultado == FALLO for f in suyas),
                sum(f.resultado in (PENDIENTE, EN_PROCESO, CANCELADO) for f in suyas),
                sum(f.documentos for f in suyas),
            ]
            for carpeta, suyas in por_carpeta.items()
        ],
        color_en=1,
    )
    tabla(
        libro.create_sheet("Por PDF"),
        [("Carpeta", 18), ("PDF", 22), ("Verificación", 22), ("Por qué", 60), ("Resultado", 18),
         ("Detalle", 60), ("Documentos", 12), ("Páginas", 10)],
        [
            [f.carpeta, f.pdf, f.verificacion, f.motivo, f.resultado, f.detalle, f.documentos, f.paginas]
            for f in filas
        ],
        color_en=4,
    )
    tabla(
        libro.create_sheet("Incidencias"),
        [("Aviso de la corrida", 100)],
        [[aviso] for aviso in avisos],
    )
    destino.mkdir(parents=True, exist_ok=True)
    ruta = destino / NOMBRE_DEL_ANALISIS
    return guardar_libro(libro, ruta)
