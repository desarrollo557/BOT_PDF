"""Leer de vuelta una planilla FUID ya escrita, para mirarla sin descargarla.

El técnico y el de calidad revisan el inventario y no se llevan el archivo de
Excel, así que hace falta enseñárselo por pantalla. Podría armarse la tabla otra
vez desde el informe del trabajo, y sería más rápido; se lee del archivo a
propósito, por lo mismo que la descarga sirve el archivo escrito y no uno
generado al vuelo: **lo que se mira y lo que se firma tienen que ser el mismo
documento.** Una vista reconstruida discrepa del disco en cuanto una de las dos
rutas cambie, y entonces revisar en pantalla deja de valer para nada.

Sólo lee. No abre la plantilla en blanco ni escribe nunca: el que rellena el
formato es `fuid_inventory.py`, y tenerlos separados es lo que impide que una
vista acabe corrigiendo la planilla que estaba enseñando.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .fuid_inventory import (
    COLUMNA_CABECERA,
    COLUMNAS,
    FILA_OBJETO,
    FILA_OFICINA,
    HOJA,
    MARCA_DE_FIRMAS,
    PRIMERA_FILA,
)

#: Las dos filas en que la plantilla escribe los encabezados de columna. El
#: formato los lleva a dos niveles -- "CONSECUTIVO" arriba, "Inicial" y "Final"
#: debajo -- y una vista que enseñara sólo uno de los dos dejaría cuatro
#: columnas llamadas "Inicial" y "Final" sin decir inicial de qué.
FILA_ENCABEZADO = 12
FILA_SUBENCABEZADO = 13

#: Cuántas filas en blanco seguidas se leen antes de dar la tabla por terminada.
#: El escritor deja dos entre la última fila y el bloque de firmas, así que con
#: una sola se cortaría la lectura de una planilla que tuviera un hueco, y sin
#: tope se leerían las firmas como si fueran registros.
VACIAS_PARA_TERMINAR = 3


@dataclass(frozen=True, slots=True)
class Grupo:
    """Un encabezado del nivel de arriba, y cuántas columnas abarca.

    El formato agrupa columnas: «CONSECUTIVO» cubre inicial y final, «UNIDAD DE
    CONSERVACION» cubre caja, carpeta, tomo y otro. Enseñar eso plano -- una
    columna llamada «CONSECUTIVO Final» y otra «FECHAS EXTREMAS Final» -- obliga
    a leer dos veces para saber final de qué, y no se parece al papel que el
    operador tiene delante.
    """

    titulo: str
    ancho: int

    def as_dict(self) -> dict[str, object]:
        return {"titulo": self.titulo, "ancho": self.ancho}


@dataclass(frozen=True, slots=True)
class PlanillaLeida:
    """Una planilla FUID tal como está en el disco."""

    archivo: str
    oficina_productora: str
    objeto: str
    #: Los nombres de columna ya aplanados, uno por columna. Es lo que hace
    #: legible una fila suelta fuera de la tabla -- un lector de pantalla, una
    #: exportación -- y por eso se conserva además de la cabecera a dos niveles.
    columnas: list[str]
    #: El nivel de arriba, con cuántas columnas ocupa cada uno.
    grupos: list[Grupo]
    #: El nivel de abajo, uno por columna y vacío donde la columna no tiene dos
    #: niveles.
    subcolumnas: list[str]
    filas: list[list[str]]

    @property
    def total(self) -> int:
        return len(self.filas)

    def as_dict(self) -> dict[str, object]:
        return {
            "archivo": self.archivo,
            "oficina_productora": self.oficina_productora,
            "objeto": self.objeto,
            "columnas": self.columnas,
            "grupos": [grupo.as_dict() for grupo in self.grupos],
            "subcolumnas": self.subcolumnas,
            "filas": self.filas,
            "total": self.total,
        }


def _texto(valor: object) -> str:
    """Una celda como la leería una persona.

    Los enteros que openpyxl devuelve como `1.0` se escriben `1`: la columna de
    folios dice cuántas hojas hay, y "1.0 folios" es una cifra que nadie escribe.
    """
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _cabecera(superiores: list[str], inferiores: list[str]) -> tuple[list[str], list[Grupo], list[str]]:
    """Los dos niveles del formato: aplanados, agrupados y por separado.

    Las celdas combinadas dejan el valor sólo en la primera del grupo, así que
    el nivel de arriba se arrastra mientras el de abajo siga diciendo algo: sin
    eso, "Final" de las fechas extremas y "Final" del consecutivo salen con el
    mismo nombre. Una columna de un solo nivel corta el arrastre, porque su
    nombre es suyo y no encabeza a la siguiente.
    """
    etiquetas: list[str] = []
    grupos: list[Grupo] = []
    subcolumnas: list[str] = []
    arriba = ""

    for superior, inferior in zip(superiores, inferiores, strict=True):
        if superior:
            arriba = superior
        if inferior:
            etiquetas.append(f"{arriba} {inferior}".strip() if arriba else inferior)
            subcolumnas.append(inferior)
            # Sigue dentro del mismo grupo mientras no aparezca uno nuevo.
            if superior or not grupos:
                grupos.append(Grupo(titulo=arriba, ancho=1))
            else:
                grupos[-1] = Grupo(titulo=grupos[-1].titulo, ancho=grupos[-1].ancho + 1)
        else:
            etiquetas.append(superior or arriba)
            subcolumnas.append("")
            grupos.append(Grupo(titulo=superior or arriba, ancho=1))
            # Una columna de un solo nivel no arrastra su nombre a la siguiente.
            arriba = ""

    return etiquetas, grupos, subcolumnas


def _fila(celdas, ancho: int = COLUMNAS) -> list[str]:
    """Una fila de la hoja como texto, rellenada hasta el ancho del formato.

    Rellenada porque openpyxl acorta la fila en la última celda con contenido, y
    una tabla cuyas filas no miden todas lo mismo desalinea la pantalla en
    cuanto una columna del final queda vacía.
    """
    valores = [_texto(celda.value) for celda in celdas][:ancho]
    return valores + [""] * (ancho - len(valores))


def leer_fuid(ruta: Path) -> PlanillaLeida:
    """Lo que dice la planilla que hay en esa ruta.

    De una sola pasada y en modo solo lectura. En ese modo openpyxl no guarda la
    hoja en memoria, así que pedir una celda por su fila y su columna vuelve a
    recorrer el archivo: sobre una planilla de mil filas serían dieciséis mil
    recorridos. `iter_rows` la lee una vez.
    """
    ruta = Path(ruta)
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = libro[HOJA] if HOJA in libro.sheetnames else libro.worksheets[0]

        superiores: list[str] = [""] * COLUMNAS
        inferiores: list[str] = [""] * COLUMNAS
        oficina = objeto = ""
        filas: list[list[str]] = []
        vacias = 0

        for numero, celdas in enumerate(hoja.iter_rows(max_col=COLUMNAS), start=1):
            valores = _fila(celdas)
            if numero == FILA_OFICINA:
                oficina = valores[COLUMNA_CABECERA - 1]
            elif numero == FILA_OBJETO:
                objeto = valores[COLUMNA_CABECERA - 1]
            elif numero == FILA_ENCABEZADO:
                superiores = valores
            elif numero == FILA_SUBENCABEZADO:
                inferiores = valores
            elif numero >= PRIMERA_FILA:
                # El bloque de firmas no es un registro. Se reconoce por su
                # texto, igual que lo hace el escritor para moverlo.
                if MARCA_DE_FIRMAS in " ".join(valores).lower():
                    break
                if not any(valores):
                    vacias += 1
                    if vacias >= VACIAS_PARA_TERMINAR:
                        break
                    continue
                vacias = 0
                filas.append(valores)

        columnas, grupos, subcolumnas = _cabecera(superiores, inferiores)
        return PlanillaLeida(
            archivo=ruta.name,
            oficina_productora=oficina,
            objeto=objeto,
            columnas=columnas,
            grupos=grupos,
            subcolumnas=subcolumnas,
            filas=filas,
        )
    finally:
        libro.close()
