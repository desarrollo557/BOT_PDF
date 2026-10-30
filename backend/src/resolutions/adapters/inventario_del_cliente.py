"""Lo que dice de una carpeta el inventario con que el cliente la entregó.

Las carpetas del REPOSITORIO UDC traen a veces, sueltas en su raíz, el Excel del
cliente: «RESOLUCIONES INTERNAS Y ACTAS.xlsx», «ACTAS DE CONSEJO DE FACULTAD DE
ODONTOLOGIA NUEVO.xlsx», una fila por UPD con su «OFICINA PRODUCTORA». Es lo que
mejor dice la cabecera del FUID cuando el nombre de la carpeta no nombra su
facultad: «18-RESOLUCIONES INTERNAS Y ACTAS 3303» es de Derecho y sólo su Excel
lo dice (ver ``domain/oficina_productora.py``).
"""

from __future__ import annotations

import logging
from collections import Counter
from pathlib import Path

logger = logging.getLogger(__name__)

#: El encabezado de la columna, tal cual en los Excel del cliente.
_CABECERA = "OFICINA PRODUCTORA"
#: Cuántas filas se miran buscando el encabezado, y cuántas columnas: los del
#: cliente declaran 16 384 columnas (A1:XFD82) y recorrerlas todas tarda.
_FILAS_DE_CABECERA = 20
_COLUMNAS = 40


def es_del_bot(nombre: str) -> bool:
    """Si el Excel lo dejó el propio bot: un general, un FUID, un inventario o un temporal."""
    return nombre.startswith(("00.", "~$")) or any(
        sufijo in nombre for sufijo in ("__FUID", "__inventario")
    )


def oficina_del_excel(ruta: Path) -> str | None:
    """La oficina productora que más se repite en un Excel del cliente, o ``None``."""
    from openpyxl import load_workbook

    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - un Excel que no se abre no dice nada
        logger.debug("no se pudo abrir %s", ruta, exc_info=True)
        return None
    cuenta: Counter[str] = Counter()
    try:
        for hoja in libro.worksheets:
            columna: int | None = None
            for numero, fila in enumerate(hoja.iter_rows(max_col=_COLUMNAS, values_only=True), start=1):
                if columna is None:
                    if numero > _FILAS_DE_CABECERA:
                        break
                    columna = next(
                        (
                            indice
                            for indice, valor in enumerate(fila)
                            if isinstance(valor, str) and " ".join(valor.split()).upper() == _CABECERA
                        ),
                        None,
                    )
                    continue
                valor = fila[columna] if columna < len(fila) else None
                if isinstance(valor, str) and valor.strip():
                    cuenta[" ".join(valor.split()).upper()] += 1
    finally:
        libro.close()
    return cuenta.most_common(1)[0][0] if cuenta else None


def _plano(valor: object) -> str:
    return " ".join(str(valor or "").split()).upper()


def folios_por_upd(ruta: Path) -> dict[str, int]:
    """Los folios -- las hojas de papel -- que el cliente anotó de cada UPD, o nada si no los trae.

    Su Excel trae «FOLIOS» (las hojas) e «IMÁGENES» (las páginas escaneadas):
    «UPD3860000» tiene 101 folios y 199 imágenes, porque casi todas sus hojas se
    escanearon por las dos caras. El bot cuenta páginas y no sabe cuántas hojas
    hay; el operador eligió el 2026-10-30 que «No. de folios» lleve las hojas del
    cliente cuando se inventaría el PDF entero.
    """
    from openpyxl import load_workbook

    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - un Excel que no se abre no dice nada
        logger.debug("no se pudo abrir %s", ruta, exc_info=True)
        return {}
    folios: dict[str, int] = {}
    try:
        for hoja in libro.worksheets:
            columnas: tuple[int, int] | None = None
            for numero, fila in enumerate(hoja.iter_rows(max_col=_COLUMNAS, values_only=True), start=1):
                if columnas is None:
                    if numero > _FILAS_DE_CABECERA:
                        break
                    rotulos = [_plano(valor) for valor in fila]
                    if "FOLIOS" in rotulos and "UPD" in rotulos:
                        columnas = (rotulos.index("FOLIOS"), rotulos.index("UPD"))
                    continue
                de_folios, de_upd = columnas
                upd = _plano(fila[de_upd] if de_upd < len(fila) else None).replace(" ", "")
                valor = fila[de_folios] if de_folios < len(fila) else None
                if upd.startswith("UPD") and isinstance(valor, int | float) and valor > 0:
                    folios[upd] = int(valor)
    finally:
        libro.close()
    return folios


def oficina_de_la_carpeta(carpeta: Path) -> str | None:
    """La oficina de los Excel del cliente sueltos en ``carpeta``, sin entrar en subcarpetas."""
    try:
        excels = sorted(
            (ruta for ruta in carpeta.glob("*.xlsx") if ruta.is_file() and not es_del_bot(ruta.name)),
            key=lambda ruta: ruta.name.lower(),
        )
    except OSError:
        return None
    for ruta in excels:
        if oficina := oficina_del_excel(ruta):
            return oficina
    return None
