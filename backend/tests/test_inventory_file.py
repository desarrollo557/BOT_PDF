"""Inventariar un PDF entero: una fila, sin partir nada.

Va de la lectura del documento a la fila del FUID y, en la última prueba, hasta
el archivo de Excel escrito sobre la plantilla oficial F-PSD-001. Lo que se
comprueba es lo que el operador pidió: que de cada PDF salga su tipo y sus
fechas extremas -- la más antigua y la más reciente que aparezcan en cualquiera
de sus hojas -- sin que el original se toque.
"""

from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

from resolutions.adapters.fuid_psd001 import HOJA, PRIMERA_FILA, FuidPsd001
from resolutions.application.archivo_fuid import SIN_ASUNTO, fila_del_archivo
from resolutions.application.fuid import NO_APLICA, Cabecera, Ubicacion
from resolutions.application.inventory_document import FORMATO_PSD001, template_named
from resolutions.application.inventory_file import InventoryFile
from resolutions.application.ports import Band
from resolutions.domain.ficha import fichar

NOTA_DE_AJUSTE = """NOTA DE AJUSTE
KINGSPAN PANELES AISLADOS S.A.S          Numero:  001-NA-00001204
NIT 900447906-1                          Fecha:   30/08/2022
Tercero: 890903938 BANCOLOMBIA S.A.
"""

EGRESO = """COMPROBANTE DE EGRESO
KINGSPAN COMERCIAL SAS                   Numero:  001-CE-0000169
NIT 901090638-1                          Fecha:   24/09/2024
"""


class FuenteFalsa:
    """Un documento ya leído, para no depender de un PDF en las pruebas."""

    def __init__(self, paginas: list[str]) -> None:
        self._paginas = paginas
        self.renderizadas: list[int] = []

    @property
    def page_count(self) -> int:
        return len(self._paginas)

    def text_of(self, page_number: int) -> str:
        return self._paginas[page_number - 1]

    def render(self, page_number: int, band: Band | None = None, dpi: int = 200) -> bytes:
        self.renderizadas.append(page_number)
        return b""

    def close(self) -> None:
        pass


class OcrFalso:
    """Un OCR que devuelve lo que se le diga, para las páginas sin capa de texto."""

    def __init__(self, texto: str = "") -> None:
        self._texto = texto
        self.veces = 0

    def read(self, image_png: bytes):
        from resolutions.application.ports import OcrResult

        self.veces += 1
        return OcrResult(text=self._texto, mean_confidence=0.9)


class TestUnaFilaPorArchivo:
    """La unidad documental es el PDF, no lo que lleve dentro."""

    def test_un_archivo_de_varias_hojas_produce_una_sola_fila(self):
        fuente = FuenteFalsa([NOTA_DE_AJUSTE, EGRESO])
        salida = InventoryFile().execute(fuente, document_name="kingspan.pdf")
        assert len(salida.rows) == 1

    def test_las_fechas_extremas_son_la_mas_antigua_y_la_mas_reciente(self):
        fuente = FuenteFalsa([NOTA_DE_AJUSTE, EGRESO])
        salida = InventoryFile().execute(fuente, document_name="kingspan.pdf")
        assert salida.row.fecha_inicial == "30-08-2022"
        assert salida.row.fecha_final == "24-09-2024"

    def test_los_folios_son_las_paginas_del_archivo(self):
        fuente = FuenteFalsa([NOTA_DE_AJUSTE, EGRESO, ""])
        salida = InventoryFile().execute(fuente, document_name="kingspan.pdf")
        assert salida.row.folios == 3

    def test_el_asunto_sale_del_encabezado_impreso(self):
        fuente = FuenteFalsa([NOTA_DE_AJUSTE])
        salida = InventoryFile().execute(fuente, document_name="kingspan.pdf")
        assert salida.row.asunto == "NOTA DE AJUSTE"


class TestLaPaginaSinCapaDeTexto:
    """Un escaneo se lee con OCR, y la página entera, no su banda superior."""

    def test_se_pasa_por_ocr_la_pagina_que_no_trae_texto(self):
        fuente = FuenteFalsa([""])
        ocr = OcrFalso(NOTA_DE_AJUSTE)
        salida = InventoryFile(ocr=ocr).execute(fuente, document_name="escaneo.pdf")
        assert ocr.veces == 1
        assert salida.ocr_pages == 1
        assert salida.row.asunto == "NOTA DE AJUSTE"

    def test_no_se_pasa_por_ocr_la_que_ya_se_puede_leer(self):
        ocr = OcrFalso("no debería hacer falta")
        salida = InventoryFile(ocr=ocr).execute(
            FuenteFalsa([NOTA_DE_AJUSTE]), document_name="nativo.pdf"
        )
        assert ocr.veces == 0
        assert salida.ocr_pages == 0

    def test_una_pagina_que_no_se_lee_de_ninguna_forma_se_declara(self):
        # Ni capa de texto ni OCR: la hoja existe, cuenta como folio, y quien
        # firme el inventario tiene derecho a saber que no se leyó.
        salida = InventoryFile(ocr=OcrFalso("")).execute(
            FuenteFalsa(["", NOTA_DE_AJUSTE]), document_name="mixto.pdf"
        )
        assert salida.unreadable == [1]
        assert salida.row.folios == 2
        assert any("no se pudieron leer" in aviso for aviso in salida.incidents)


class TestLoQueNoSeLeeSeDeclara:
    """Nada se rellena con algo verosímil."""

    def test_un_archivo_ilegible_lo_dice_en_el_asunto_y_en_las_notas(self):
        salida = InventoryFile().execute(FuenteFalsa(["   "]), document_name="vacio.pdf")
        assert salida.row.asunto == SIN_ASUNTO
        assert salida.row.fecha_inicial == NO_APLICA
        assert salida.row.fecha_final == NO_APLICA
        assert "no se pudo leer el encabezado" in salida.row.notas

    def test_la_falta_de_consecutivo_no_se_advierte(self):
        # El operador lo dijo de estos papeles: no todos traen número impreso.
        # Advertirlo en cada fila llenaría el inventario de una nota que no pide
        # ninguna acción, y enterraría las que sí.
        ficha = fichar(["ACTA DE ENTREGA\nFecha: 12/05/2023"])
        fila = fila_del_archivo(ficha, nombre_del_archivo="acta.pdf")
        assert fila.consecutivo_inicial == NO_APLICA
        assert "número" not in fila.notas


class TestElFuidQueQuedaEscrito:
    """La fila cae en la columna que le toca de la plantilla oficial."""

    @pytest.fixture
    def escrito(self, tmp_path):
        plantilla = template_named(FORMATO_PSD001)
        assert plantilla is not None and plantilla.is_file()
        ficha = fichar([NOTA_DE_AJUSTE, EGRESO], folios=2)
        fila = fila_del_archivo(
            ficha, nombre_del_archivo="kingspan.pdf", ubicacion=Ubicacion(caja="14")
        )
        destino = tmp_path / "salida.xlsx"
        FuidPsd001(plantilla).write(
            [fila],
            destino,
            cabecera=Cabecera(entidad_productora="KINGSPAN PANELES AISLADOS S.A.S"),
            inventariado=date(2026, 9, 15),
        )
        return load_workbook(destino)[HOJA]

    def test_el_asunto_y_las_fechas_extremas_van_donde_dice_el_formato(self, escrito):
        # J es ASUNTOS, M y N son las dos mitades de FECHAS EXTREMAS.
        assert escrito.cell(row=PRIMERA_FILA, column=10).value == "NOTA DE AJUSTE"
        assert escrito.cell(row=PRIMERA_FILA, column=13).value == "30-08-2022"
        assert escrito.cell(row=PRIMERA_FILA, column=14).value == "24-09-2024"

    def test_los_folios_y_la_caja_tambien(self, escrito):
        assert escrito.cell(row=PRIMERA_FILA, column=20).value == 2
        assert escrito.cell(row=PRIMERA_FILA, column=15).value == "14"

    def test_las_columnas_que_el_sistema_no_puede_saber_salen_como_na(self, escrito):
        # El instructivo del formato: lo que no aplica se escribe N/A. Una celda
        # vacía no dice "no aplica", dice "se olvidó".
        for columna in (3, 8, 9, 19, 26, 27):
            assert escrito.cell(row=PRIMERA_FILA, column=columna).value == NO_APLICA

    def test_la_cabecera_de_la_entrega_se_repite_en_la_fila(self, escrito):
        assert escrito.cell(row=PRIMERA_FILA, column=4).value == (
            "KINGSPAN PANELES AISLADOS S.A.S"
        )

    def test_los_encabezados_de_la_plantilla_no_se_tocan(self, escrito):
        assert escrito.cell(row=7, column=10).value == "ASUNTOS"
        assert escrito.cell(row=7, column=13).value == "Inicial"
        assert escrito.cell(row=7, column=14).value == "Final"
