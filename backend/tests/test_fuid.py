"""Correspondencia con el Formato Único de Inventario Documental.

Lo que se comprueba aquí no es que el Excel se escriba, sino que se respeten
las dos reglas del instructivo del formato que es fácil incumplir sin darse
cuenta: lo que no aplica se escribe N/A, y una fecha que no se pudo leer no se
sustituye por una plausible.
"""

from dataclasses import replace  # noqa: E402

import pytest
from openpyxl import Workbook, load_workbook

from resolutions.adapters.fuid_inventory import (
    BLOQUE_FIRMAS,
    FILAS_ANTES_DE_FIRMAR,
    PRIMERA_FILA,
    FuidInventory,
)
from resolutions.application.diploma_fuid import (
    asunto,
    fecha_fuid,
    filas_del_libro,
    folio_confiable,
)
from resolutions.application.diploma_split import group_by_record
from resolutions.application.fuid import (
    NO_APLICA,
    Cabecera,
    FuidRow,
    Ubicacion,
    upd_del_archivo,
)
from resolutions.domain.diploma import DiplomaRecord

LIBRO = "6.3858084 REGISTRO DE DIPLOMAS N°08 2012-2013.pdf"

REGISTRO = DiplomaRecord(
    page_number=1,
    folio="336BIS",
    registered_folio="336BIS",
    book="8",
    name="ALIX JOSEFINA MARIN RODRIGUEZ",
    identity_kind="C.C",
    identity_number="22793650",
    degree="ESPECIALISTA EN GESTION DE LA CALIDAD Y AUDITORIA EN SALUD",
    graduation_date="29/03/2012",
)


class TestFechas:
    def test_convierte_al_formato_del_encabezado(self):
        assert fecha_fuid("29/03/2012") == "29-03-2012"

    def test_rellena_el_dia_y_el_mes_a_dos_cifras(self):
        assert fecha_fuid("3/8/2012") == "03-08-2012"

    def test_una_fecha_ausente_se_escribe_n_a(self):
        assert fecha_fuid(None) == NO_APLICA

    def test_una_fecha_ilegible_no_se_adivina(self):
        assert fecha_fuid("3 de agosto") == NO_APLICA


class TestUnidadDocumental:
    def test_lee_el_upd_del_nombre_del_archivo(self):
        assert upd_del_archivo(LIBRO) == "3858084"

    def test_lo_escribe_pelado_sin_anteponerle_upd(self):
        assert not upd_del_archivo(LIBRO).startswith("UPD")

    def test_lo_toma_tambien_de_los_libros_de_posgrados(self):
        nombre = "4. 3858082 REGISTRO DE DIPLOMAS POSGRADOS N°08 18-15-21-80.pdf"
        assert upd_del_archivo(nombre) == "3858082"

    def test_sin_upd_en_el_nombre_se_escribe_n_a(self):
        assert upd_del_archivo("libro escaneado.pdf") == NO_APLICA


class TestAsunto:
    def test_junta_el_nombre_y_el_titulo(self):
        assert asunto(REGISTRO).startswith("ALIX JOSEFINA MARIN RODRIGUEZ - ESPECIALISTA")

    def test_con_un_solo_dato_entrega_ese_dato(self):
        assert asunto(DiplomaRecord(page_number=1, name="JUANA DE LA CRUZ")) == (
            "JUANA DE LA CRUZ"
        )

    def test_sin_ninguno_de_los_dos_escribe_n_a(self):
        assert asunto(DiplomaRecord(page_number=1)) == NO_APLICA


class TestFolio:
    def test_avisa_de_la_discrepancia(self):
        folio, discrepan = folio_confiable(
            DiplomaRecord(page_number=213, folio="5", registered_folio="543")
        )
        assert (folio, discrepan) == ("543", True)

    def test_sin_folio_legible_escribe_n_a(self):
        assert folio_confiable(DiplomaRecord(page_number=152))[0] == NO_APLICA


class TestFilas:
    def test_una_fila_por_pagina_en_el_orden_del_pdf(self):
        registros = [DiplomaRecord(page_number=n, name=f"NOMBRE {n}") for n in (1, 2, 3)]
        filas = filas_del_libro(registros, nombre_del_archivo=LIBRO)
        assert [f.orden for f in filas] == [1, 2, 3]

    def test_la_numeracion_continua_entre_libros(self):
        registros = [DiplomaRecord(page_number=1, name="ALGUIEN CONOCIDO")]
        filas = filas_del_libro(registros, nombre_del_archivo=LIBRO, desde=398)
        assert filas[0].orden == 398

    def test_el_folio_de_registro_va_en_notas(self):
        fila = filas_del_libro([REGISTRO], nombre_del_archivo=LIBRO)[0]
        assert fila.notas == "Registrado a folio No. 336BIS"

    def test_el_libro_empastado_va_en_tomo(self):
        assert filas_del_libro([REGISTRO], nombre_del_archivo=LIBRO)[0].tomo == "8"

    def test_la_cedula_va_en_el_consecutivo_inicial(self):
        fila = filas_del_libro([REGISTRO], nombre_del_archivo=LIBRO)[0]
        assert fila.consecutivo_inicial == "22793650"

    def test_las_dos_fechas_extremas_son_la_del_grado(self):
        fila = filas_del_libro([REGISTRO], nombre_del_archivo=LIBRO)[0]
        assert fila.fecha_inicial == fila.fecha_final == "29-03-2012"

    def test_cada_registro_ocupa_un_folio(self):
        fila = filas_del_libro([REGISTRO], nombre_del_archivo=LIBRO)[0]
        assert (fila.folios, fila.folios_siar) == (1, 1)

    def test_la_caja_no_se_inventa(self):
        fila = filas_del_libro([REGISTRO], nombre_del_archivo=LIBRO)[0]
        assert fila.caja == NO_APLICA

    def test_la_ubicacion_se_usa_cuando_la_dan(self):
        fila = filas_del_libro(
            [REGISTRO],
            nombre_del_archivo=LIBRO,
            ubicacion=Ubicacion(caja="050C000001", otro="1018-5", codigo_trd="140-25"),
        )[0]
        assert (fila.caja, fila.otro, fila.codigo_trd) == ("050C000001", "1018-5", "140-25")

    def test_una_pagina_sin_datos_no_deja_celdas_en_blanco(self):
        fila = filas_del_libro([DiplomaRecord(page_number=152)], nombre_del_archivo=LIBRO)[0]
        celdas = fila.as_cells()
        # El número de orden, los dos contadores de folios y las notas pueden no
        # ser N/A; todo lo demás sí, porque no hay nada que escribir.
        assert celdas[1] == NO_APLICA and celdas[2] == NO_APLICA
        assert celdas[5] == celdas[6] == NO_APLICA


# -----------------------------------------------------------------------------
#  El escritor, contra una plantilla mínima con la misma estructura que la real
# -----------------------------------------------------------------------------


@pytest.fixture
def plantilla(tmp_path):
    """Una plantilla con lo que el escritor necesita: encabezados, área de datos
    y el bloque de firmas en las filas 109 a 112."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "INVENTARIO"
    hoja.cell(row=8, column=2, value="OFICINA PRODUCTORA: FACULTAD DE INGENIERIA")
    hoja.cell(row=9, column=2, value="OBJETO: EL QUE VENGA EN LA PLANTILLA")
    hoja.cell(row=12, column=1, value="No. de Orden")
    hoja.cell(row=PRIMERA_FILA, column=1, value="consecutivo desde 1")
    hoja.cell(row=PRIMERA_FILA, column=3, value="mobre- titulo")

    inicio = BLOQUE_FIRMAS[0]
    hoja.cell(row=inicio, column=1, value="Elaborado por:")
    hoja.cell(row=inicio, column=3, value="_______________")
    hoja.merge_cells(start_row=inicio, start_column=3, end_row=inicio, end_column=4)
    hoja.cell(row=inicio + 3, column=1, value="Lugar:")

    ruta = tmp_path / "plantilla.xlsx"
    libro.save(ruta)
    return ruta


def escribir(plantilla, tmp_path, filas, cabecera=None):
    destino = tmp_path / "inventario.xlsx"
    FuidInventory(plantilla).write(filas, destino, cabecera=cabecera)
    return load_workbook(destino)["INVENTARIO"]


def fila_ejemplo(orden):
    return FuidRow(orden=orden, asunto=f"GRADUANDO {orden} - ESPECIALISTA", tomo="8")


class TestEscritura:
    def test_el_primer_registro_sustituye_la_fila_guia(self, plantilla, tmp_path):
        hoja = escribir(plantilla, tmp_path, [fila_ejemplo(1)])
        assert hoja.cell(row=14, column=1).value == 1
        assert hoja.cell(row=14, column=3).value == "GRADUANDO 1 - ESPECIALISTA"

    def test_baja_las_firmas_cuando_los_registros_no_caben(self, plantilla, tmp_path):
        filas = [fila_ejemplo(n) for n in range(1, 201)]
        hoja = escribir(plantilla, tmp_path, filas)

        ultima = PRIMERA_FILA + len(filas) - 1
        assert hoja.cell(row=ultima, column=1).value == 200
        firma = ultima + FILAS_ANTES_DE_FIRMAR + 1
        assert hoja.cell(row=firma, column=1).value == "Elaborado por:"
        assert hoja.cell(row=firma + 3, column=1).value == "Lugar:"

    def test_conserva_la_combinacion_de_celdas_de_la_firma(self, plantilla, tmp_path):
        filas = [fila_ejemplo(n) for n in range(1, 21)]
        hoja = escribir(plantilla, tmp_path, filas)
        firma = PRIMERA_FILA + len(filas) - 1 + FILAS_ANTES_DE_FIRMAR + 1
        assert any(
            rango.min_row == firma and rango.min_col == 3 and rango.max_col == 4
            for rango in hoja.merged_cells.ranges
        )

    def test_no_deja_rastro_de_las_firmas_en_su_sitio_anterior(self, plantilla, tmp_path):
        hoja = escribir(plantilla, tmp_path, [fila_ejemplo(1)])
        assert hoja.cell(row=BLOQUE_FIRMAS[0], column=1).value is None

    def test_cambia_la_oficina_productora_cuando_se_indica(self, plantilla, tmp_path):
        hoja = escribir(
            plantilla,
            tmp_path,
            [fila_ejemplo(1)],
            cabecera=Cabecera(oficina_productora="SECRETARÍA GENERAL"),
        )
        assert hoja.cell(row=8, column=2).value == "OFICINA PRODUCTORA: SECRETARÍA GENERAL"

    def test_sin_cabecera_no_toca_lo_que_traiga_la_plantilla(self, plantilla, tmp_path):
        hoja = escribir(plantilla, tmp_path, [fila_ejemplo(1)])
        assert hoja.cell(row=9, column=2).value == "OBJETO: EL QUE VENGA EN LA PLANTILLA"


class TestLaVueltaDeUnFolioNoEsUnaFila:
    """El inventario y la carpeta tienen que contar lo mismo.

    Una página sin ningún identificador es la vuelta de la hoja anterior -- el
    archivo escanea por las dos caras y lo anota con una "v" junto al folio
    escrito a mano -- o algo anexado. El separador la mete en el PDF del
    registro anterior; si el FUID le abriera fila propia, prometería más
    diplomas de los que hay archivos, y con el asunto en blanco.
    """

    def vacia(self, pagina):
        return DiplomaRecord(page_number=pagina)

    def test_it_does_not_open_a_row_of_its_own(self):
        filas = filas_del_libro(
            [REGISTRO, self.vacia(REGISTRO.page_number + 1)], nombre_del_archivo=LIBRO
        )
        assert len(filas) == 1

    def test_it_counts_as_a_folio_of_the_row_above(self):
        filas = filas_del_libro(
            [REGISTRO, self.vacia(REGISTRO.page_number + 1)], nombre_del_archivo=LIBRO
        )
        assert filas[0].folios == 2
        assert filas[0].folios_siar == 2

    def test_the_order_column_stays_consecutive(self):
        # "Debe anotarse en forma consecutiva el número correspondiente de cada
        # unidad documental", dice el instructivo. Saltarse un número porque una
        # página no era una unidad rompería justamente eso.
        registros = [
            REGISTRO,
            self.vacia(90),
            replace(REGISTRO, page_number=91, folio="341"),
            self.vacia(92),
            replace(REGISTRO, page_number=93, folio="342"),
        ]
        filas = filas_del_libro(registros, nombre_del_archivo=LIBRO)
        assert [fila.orden for fila in filas] == [1, 2, 3]

    def test_the_inventory_and_the_split_agree_on_how_many_units_there_are(self):
        registros = [
            REGISTRO,
            self.vacia(90),
            replace(REGISTRO, page_number=91, folio="341"),
        ]
        filas = filas_del_libro(registros, nombre_del_archivo=LIBRO)
        grupos = group_by_record(registros).groups
        assert len(filas) == len(grupos) == 2
        assert [fila.folios for fila in filas] == [group.size for group in grupos]
