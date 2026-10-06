"""Inventariar actas de consejo: una fila por PDF, con lo que el operador dictó.

El ejemplo que dio el operador es el acta del Consejo de Facultad de Odontología
del 17 de diciembre de 2012 (`1N°22-17 DE DICIEMBRE DE 2012.pdf`, 22 páginas):
de ella sale «ACTA ORDINARIA No. 022», «Lunes 17 de diciembre de 2012», los
folios 1 a 21 escritos a mano en la esquina, cuántos folios son y un resumen de
todo el documento que no menciona el llamado a lista ni el quórum. Todo en un
Excel sencillo, no en el FUID.

Ninguna prueba de aquí habla con un proveedor: las lecturas de pago van con
dobles y el HTTP, simulado. Las pruebas no se pagan.
"""

from __future__ import annotations

import io
import json
import urllib.error
from dataclasses import asdict
from datetime import date
from pathlib import Path

import pytest

from resolutions.adapters.esquinas_del_folio import LectorDeFolios
from resolutions.adapters.mistral_resumen import (
    INSTRUCCIONES,
    PROVEEDOR,
    MistralResumenConfig,
    MistralResumidor,
)
from resolutions.application.acta_consejo import FilaDeActa, InventariarActa
from resolutions.application.consumo import Consumo
from resolutions.application.ports import OcrResult, ResumenNoDisponible
from resolutions.application.tipo_pedido import TipoPedido, tipos_declarables
from resolutions.domain.acta_consejo import (
    TramoDeActa,
    actas_del_pdf,
    depurar_resumen,
    folios_por_continuidad,
    folios_que_cuadran,
    leer_acta,
    numeros_escritos,
    sin_llamado_a_lista,
    texto_para_resumir,
)

# Lo que traen de verdad las páginas del acta de ejemplo, recortado.
PORTADA = ",/■ '■?\n■;\nNo. 22 -17 DE DICIEMBRE DE\n2012\nj."
CITACION = (
    "Universidad\nde Cartagena\nCONSEJO DE FACULTAD\nSESION ORDINARIA\n"
    "FECHA: LUNES 17 DE DICIEMBRE DE 2012\nLUGAR: DECANATURA\nHORA: 09:30 HORAS.\n"
    "ORDEN DEL DIA:\n1. LLAMADO A LISTA Y VERIFICACION DEL QUORUM\n"
    "2. APROBACIÓN ACTA ANTERIOR\n3. ASUNTOS ACADEMICOS\n4. PROPOSICIONES O VARIOS\n"
    "Facultad de Odontología - Secretaria Decanatura"
)


def hoja(texto: str) -> str:
    """Una hoja del acta, con el membrete y el pie que llevan todas."""
    return (
        f"Universidad\nde Cartagena\nFundada en 1827\n{texto}\n"
        "Facultad de Odontología - Departamento Académico"
    )


ENCABEZADO = hoja(
    "ACTA ORDINARIA No. 022 \nLunes 17 de diciembre de 2012\nCONSEJO DE FACULTAD\n"
    "Fecha\n17 de diciembre 2012\nDra. LUISA AREVALO TOVAR- DECANA"
)
CUERPO = hoja(
    "Acta No. 022-2012 \nConsejo de Facultad\n"
    "1, LLAMADO A LISTA Y VERIFICACIÓN DEL\nQUÓRUM\n"
    "La señora decana, la doctora LUISA LEONOR \nAREVALO TOVAR, da inicio a la sesión, ordenando \n"
    "llamado a lista y la verificación del quórum. Se verifica el \n"
    "quórum, suficiente para sesionar y decidir. Se lee y se \n"
    "aprueba el acta anterior. Se aprueba el orden del día.\n"
    "Consejo Académico de fecha 13 de Diciembre de 2012.\n"
    "Según consta en el\nActa No. 021 se aprobó el calendario."
)
SIGUE = hoja(
    "Acta No. 022-2012 \nINTERVENCION DE LA DECANA\n"
    "La decana presentó el informe del Consejo Académico y la reacreditación del programa."
)
CIERRE = hoja("Siendo las 11:30 a.m., finaliza la sesión.\nLa presente acta consta de 17 folios")
ACTA = [PORTADA, CITACION, ENCABEZADO, CUERPO, SIGUE, CIERRE]


# -----------------------------------------------------------------------------
#  Lo que el acta dice de sí misma
# -----------------------------------------------------------------------------


class TestLeerElActa:
    def test_el_ejemplo_del_operador(self):
        acta = leer_acta(ACTA)

        assert acta.numero == "022"
        assert acta.sesion == "ORDINARIA"
        assert acta.fecha == date(2012, 12, 17)
        assert acta.pagina == 3
        assert acta.encabezado == "ACTA ORDINARIA No. 022"
        assert acta.fecha_larga == "Lunes 17 de diciembre de 2012"

    def test_el_numero_conserva_sus_ceros(self):
        assert leer_acta(["ACTA ORDINARIA No. 007\n3 de marzo de 2010"]).numero == "007"

    def test_una_cita_en_mitad_de_la_frase_no_es_el_encabezado(self):
        # «según consta en el Acta No. 021»: es otra acta, y su número en la
        # fila sería una mentira que nadie podría rastrear.
        acta = leer_acta(["El consejo aprobó lo que consta en el Acta No. 021 de la sesión."])

        assert acta.numero is None
        assert acta.encabezado is None

    def test_el_encabezado_con_la_sesion_manda_sobre_el_membrete_corrido(self):
        acta = leer_acta(["Acta No. 022-2012\nConsejo", "ACTA EXTRAORDINARIA N° 5\n2 de mayo de 2011"])

        assert acta.numero == "5"
        assert acta.sesion == "EXTRAORDINARIA"

    def test_la_sesion_se_toma_de_la_citacion_si_el_encabezado_no_la_dice(self):
        acta = leer_acta(["SESIÓN ORDINARIA\nFECHA: 3 de marzo de 2010", "ACTA No. 14\n"])

        assert acta.encabezado == "ACTA ORDINARIA No. 14"

    def test_sin_fecha_tras_el_encabezado_vale_la_rotulada(self):
        acta = leer_acta(["FECHA: LUNES 17 DE DICIEMBRE DE 2012", "ACTA ORDINARIA No. 022\nCONSEJO"])

        assert acta.fecha == date(2012, 12, 17)

    def test_la_fecha_en_cifras_tras_el_rotulo(self):
        acta = leer_acta(["ACTA No. 3\nFecha: 05/03/2015\nLugar: Rectoría"])

        assert acta.fecha == date(2015, 3, 5)

    def test_el_dia_de_la_semana_sale_del_calendario(self):
        # El papel no lo dice y la fila lo lleva igual, sin inventar nada.
        acta = leer_acta(["ACTA ORDINARIA No. 9\n17 de diciembre de 2012"])

        assert acta.fecha_larga == "Lunes 17 de diciembre de 2012"
        assert not acta.dia_no_cuadra

    def test_un_dia_de_la_semana_que_no_cuadra_se_advierte(self):
        # Un 17 leído como 11 cae en martes, y el papel decía lunes.
        acta = leer_acta(["ACTA ORDINARIA No. 9\nLunes 11 de diciembre de 2012"])

        assert acta.dia_no_cuadra

    def test_las_tildes_que_el_ocr_se_come_no_importan(self):
        acta = leer_acta(["ACTA ORDINARIA No. 4\nMiercoles 3 de febrero de 2016"])

        assert acta.fecha == date(2016, 2, 3)
        assert acta.fecha_larga == "Miércoles 3 de febrero de 2016"

    def test_una_fecha_del_futuro_no_es_la_de_la_sesion(self):
        acta = leer_acta(["ACTA No. 1\n3 de marzo de 2099"], hoy=date(2026, 10, 2))

        assert acta.fecha is None


class TestLoQueSeMandaAResumir:
    def test_no_lleva_el_llamado_a_lista_ni_el_quorum(self):
        texto = texto_para_resumir(ACTA)

        plano = texto.lower()
        assert "llamado a lista" not in plano
        assert "quórum" not in plano and "quorum" not in plano
        # Y lo que viene detrás sí sigue.
        assert "aprueba el acta anterior" in texto
        assert "2. APROBACIÓN ACTA ANTERIOR" in texto

    def test_quita_el_membrete_que_se_repite_en_todas_las_hojas(self):
        texto = texto_para_resumir(ACTA)

        assert "Fundada en 1827" not in texto
        assert "Facultad de Odontología - Departamento Académico" not in texto

    def test_quita_el_ruido_de_los_sellos(self):
        texto = texto_para_resumir(["Se aprobó el presupuesto.\n■;\nf>C-C.F.RlS.3470\n4Í»"])

        assert texto == "Se aprobó el presupuesto."

    def test_tiene_techo(self):
        texto = texto_para_resumir(["palabra " * 20_000], limite=1_000)

        assert len(texto) <= 1_010
        assert "[…]" in texto

    def test_la_lista_numerada_no_se_lleva_el_orden_del_dia(self):
        limpio = sin_llamado_a_lista("1. LLAMADO A LISTA Y VERIFICACION DEL QUORUM\n2. INFORMES")

        assert "2. INFORMES" in limpio


class TestDepurarElResumen:
    def test_quita_la_frase_que_menciona_el_quorum(self):
        resumen = depurar_resumen(
            "Sesionó el Consejo de Facultad. Se verificó el quórum reglamentario. "
            "Se aprobó el calendario académico."
        )

        assert resumen == "Sesionó el Consejo de Facultad. Se aprobó el calendario académico."

    def test_y_la_del_llamado_a_lista(self):
        assert "lista" not in depurar_resumen("Se hizo el llamado a lista. Se trató X.")


# -----------------------------------------------------------------------------
#  El folio de la esquina
# -----------------------------------------------------------------------------


class TestLosFolios:
    def test_lo_que_leyo_el_proveedor_sobre_el_acta_de_ejemplo(self):
        # Respuesta real de Mistral OCR al mosaico de las hojas 2 y 22.
        assert folios_que_cuadran(numeros_escritos("1\n\n21"), 21) == (1, 21)

    def test_un_numero_del_membrete_no_descuadra_la_lectura(self):
        # «Fundada en 1827» no pasa por folio, y un «27» o un «5» sueltos no
        # están a veinte de ningún otro.
        numeros = numeros_escritos("Fundada en 1827\n27 1\nO/AR-005\n21")

        assert folios_que_cuadran(numeros, 21) == (1, 21)

    def test_los_codigos_y_las_fechas_no_son_folios(self):
        assert numeros_escritos("FO-DO/AR-005 FECHA: 04-05-2011") == []

    def test_dos_parejas_posibles_no_se_adivinan(self):
        assert folios_que_cuadran([1, 2, 21, 22], 21) is None

    def test_sin_numeros_no_hay_folios(self):
        assert folios_que_cuadran([], 21) is None

    def test_un_acta_de_una_sola_hoja_foliada(self):
        assert folios_que_cuadran([7], 1) == (7, 7)

    def test_la_foliacion_real_no_cuadra_siempre_exacta_y_se_le_perdona(self):
        # Carpeta UPD3859771: 42 hojas foliadas del 194 al 236, 40 del 43 al 84.
        assert folios_que_cuadran([194, 236], 42) == (194, 236)
        assert folios_que_cuadran([43, 84], 40) == (43, 84)

    def test_las_lecturas_malas_de_esa_carpeta_no_pasan(self):
        assert folios_que_cuadran([381, 421], 20) is None
        assert folios_que_cuadran([18, 14], 15) is None
        assert folios_que_cuadran([142, 123], 32) is None


class TestLaContinuidadDeLaCarpeta:
    def test_el_hueco_entre_dos_actas_leidas_se_llena_si_cuadra_exacto(self):
        # Lo que pasó en UPD3859771: la 6 acaba en 126, la 10 empieza en 194, y
        # las tres del medio tienen 15, 32 y 20 hojas: 67, el hueco justo.
        actas = [(122, 126, 5), (None, None, 15), (None, None, 32), (None, None, 20), (194, 236, 42)]

        assert folios_por_continuidad(actas) == [
            (122, 126, False),
            (127, 141, True),
            (142, 173, True),
            (174, 193, True),
            (194, 236, False),
        ]

    def test_si_no_cuadra_exacto_no_se_deduce_nada(self):
        actas = [(1, 21, 21), (None, None, 15), (40, 60, 21)]

        assert folios_por_continuidad(actas)[1] == (None, None, False)

    def test_sin_una_lectura_a_cada_lado_tampoco(self):
        actas = [(None, None, 10), (11, 20, 10)]

        assert folios_por_continuidad(actas)[0] == (None, None, False)

    def test_la_carpeta_va_en_orden_natural_y_dice_lo_deducido(self):
        from resolutions.application.acta_consejo import inventario_de_la_carpeta

        def fila(archivo, inicial, final, folios, observaciones=""):
            return {
                **_fila("ACTA", archivo),
                "folio_inicial": inicial,
                "folio_final": final,
                "folios": folios,
                "observaciones": observaciones,
            }

        filas = inventario_de_la_carpeta(
            [
                fila("10N°003.pdf", 37, 50, 14),
                fila("1N°22.pdf", 1, 21, 21),
                fila("2N°20.pdf", "N/A", "N/A", 15, "sin folio inicial y final: no cuadra"),
            ]
        )

        assert [f["archivo"] for f in filas] == ["1N°22.pdf", "2N°20.pdf", "10N°003.pdf"]
        assert (filas[1]["folio_inicial"], filas[1]["folio_final"]) == (22, 36)
        assert "folios deducidos" in filas[1]["observaciones"]
        assert "sin folio inicial y final" not in filas[1]["observaciones"]


# -----------------------------------------------------------------------------
#  El caso de uso, con dobles
# -----------------------------------------------------------------------------


class Fuente:
    """Un acta de mentira: texto por página, y tinta en la esquina de las foliadas."""

    def __init__(self, paginas, foliadas=None):
        self._paginas = paginas
        self._foliadas = set(foliadas if foliadas is not None else range(2, len(paginas) + 1))

    @property
    def page_count(self):
        return len(self._paginas)

    def text_of(self, numero):
        return self._paginas[numero - 1]

    def render(self, numero, band=None, dpi=200):
        from PIL import Image

        salida = io.BytesIO()
        Image.new("L", (40, 20), 255).save(salida, "PNG")
        return salida.getvalue()

    def ink_of(self, numero, band):
        return (0.02, 0.0) if numero in self._foliadas else (0.0, 0.0)


class OcrDeCasa:
    def __init__(self, texto="ACTA ORDINARIA No. 022\nLunes 17 de diciembre de 2012"):
        self.texto = texto
        self.llamadas = 0

    def read(self, imagen):
        self.llamadas += 1
        return OcrResult(text=self.texto)


class Resumidor:
    def __init__(self, resumen="Sesionó el Consejo de Facultad y aprobó el calendario."):
        self.resumen = resumen
        self.recibido: str | None = None

    def resumir(self, texto):
        self.recibido = texto
        return self.resumen


class Folios:
    def __init__(self, texto="1\n\n21"):
        self.texto = texto
        self.pedidas: list[list[int]] = []

    def leer_esquinas(self, source, paginas):
        self.pedidas.append(list(paginas))
        return self.texto


def acta_de_22_paginas():
    """La forma del ejemplo: portada sin folio y veintiuna hojas foliadas."""
    return Fuente([PORTADA, CITACION, ENCABEZADO, CUERPO, SIGUE] + [CIERRE] * 17)


class TestInventariarUnActa:
    def test_la_fila_lleva_lo_que_dicto_el_operador(self):
        folios = Folios()
        outcome = InventariarActa(
            ocr=OcrDeCasa(), resumidor=Resumidor(), lector_de_folios=folios
        ).execute(acta_de_22_paginas(), document_name="1N°22-17 DE DICIEMBRE DE 2012.pdf")

        fila = outcome.fila
        assert fila.archivo == "1N°22-17 DE DICIEMBRE DE 2012.pdf"
        assert fila.acta == "ACTA ORDINARIA No. 022"
        assert fila.fecha == "17-12-2012"
        assert (fila.folio_inicial, fila.folio_final, fila.folios) == (1, 21, 21)
        assert fila.resumen == "Sesionó el Consejo de Facultad y aprobó el calendario."
        assert fila.observaciones == ""
        # Sólo la primera y la última hoja foliadas, en una lectura.
        assert folios.pedidas == [[2, 22]]

    def test_a_la_ia_le_llega_texto_limpio_y_no_la_imagen(self):
        resumidor = Resumidor()
        InventariarActa(ocr=OcrDeCasa(), resumidor=resumidor).execute(
            acta_de_22_paginas(), document_name="acta.pdf"
        )

        assert isinstance(resumidor.recibido, str)
        assert "quórum" not in resumidor.recibido.lower()
        assert "llamado a lista" not in resumidor.recibido.lower()

    def test_el_resumen_es_de_todo_el_documento(self):
        # Lo que se trató va de la primera hoja a la última: el texto que se
        # manda a resumir trae la citación, el acta y lo que viene detrás.
        resumidor = Resumidor()
        InventariarActa(resumidor=resumidor).execute(
            Fuente([CITACION, ENCABEZADO, SIGUE]), document_name="a.pdf"
        )

        assert "ASUNTOS ACADEMICOS" in resumidor.recibido
        assert "ACTA ORDINARIA No. 022" in resumidor.recibido
        assert "reacreditación del programa" in resumidor.recibido

    def test_el_ocr_de_casa_solo_lee_la_pagina_sin_capa(self):
        # La portada trae 46 caracteres de capa: va al OCR. Las demás, no.
        ocr = OcrDeCasa()
        outcome = InventariarActa(ocr=ocr).execute(acta_de_22_paginas(), document_name="a.pdf")

        assert ocr.llamadas == 1
        assert outcome.ocr_pages == 1

    def test_con_solo_ocr_se_inventaria_igual_y_se_dice_que_falto(self):
        outcome = InventariarActa(
            ocr=OcrDeCasa(), motivo_sin_resumidor="se leyó con «Solo OCR»"
        ).execute(acta_de_22_paginas(), document_name="a.pdf")

        fila = outcome.fila
        assert fila.acta == "ACTA ORDINARIA No. 022"
        assert fila.fecha == "17-12-2012"
        # Los folios contados por la tinta, sin leer ninguno.
        assert fila.folios == 21
        assert fila.folio_inicial == fila.folio_final == "N/A"
        assert fila.resumen == "N/A"
        # Lo que falta por la elección del operador no se repite en cada fila:
        # lo dice una vez la cabecera del Excel.
        assert fila.observaciones == ""

    def test_con_solo_ocr_la_cabecera_del_excel_lo_dice_una_vez(self, tmp_path):
        from openpyxl import load_workbook

        from resolutions.adapters.excel_actas import escribir_inventario_de_actas
        from resolutions.application.acta_consejo import SIN_IA

        outcome = InventariarActa(
            ocr=OcrDeCasa(), motivo_sin_resumidor="se leyó con «Solo OCR»"
        ).execute(acta_de_22_paginas(), document_name="a.pdf")
        destino = escribir_inventario_de_actas(outcome.filas, tmp_path / "a.xlsx", nota=SIN_IA)

        hoja = load_workbook(destino).active
        assert "Solo OCR" in hoja["A2"].value
        assert _indicador(hoja, "POR REVISAR") == 0

    def test_un_acta_sin_folios_no_gasta_en_leerlos(self):
        folios = Folios()
        fuente = Fuente([ENCABEZADO, CUERPO, SIGUE], foliadas=[])
        outcome = InventariarActa(lector_de_folios=folios).execute(fuente, document_name="a.pdf")

        assert folios.pedidas == []
        assert outcome.fila.folios == 3
        assert "ninguna hoja lleva folio" in outcome.fila.observaciones

    def test_una_lectura_que_no_cuadra_deja_los_folios_en_na_y_lo_dice(self):
        outcome = InventariarActa(lector_de_folios=Folios("4\n\n70")).execute(
            acta_de_22_paginas(), document_name="a.pdf"
        )

        assert outcome.fila.folio_inicial == "N/A"
        assert outcome.fila.folios == 21
        assert "no cuadra con las 21 hojas foliadas" in outcome.fila.observaciones

    def test_un_resumen_que_no_llega_no_tumba_el_acta(self):
        class SinSaldo:
            def resumir(self, texto):
                raise ResumenNoDisponible("la cuenta de Mistral no tiene saldo (HTTP 402)")

        outcome = InventariarActa(resumidor=SinSaldo()).execute(
            acta_de_22_paginas(), document_name="a.pdf"
        )

        assert outcome.resumen is None
        assert "sin resumen: la cuenta de Mistral no tiene saldo" in outcome.fila.observaciones

    def test_el_resumen_se_depura_aunque_el_modelo_desobedezca(self):
        outcome = InventariarActa(
            resumidor=Resumidor("Se verificó el quórum. Se aprobó el calendario.")
        ).execute(acta_de_22_paginas(), document_name="a.pdf")

        assert outcome.fila.resumen == "Se aprobó el calendario."

    def test_sin_numero_ni_fecha_la_fila_lo_declara(self):
        outcome = InventariarActa().execute(
            Fuente(["Texto cualquiera sin encabezado de acta, bastante largo " * 3]),
            document_name="a.pdf",
        )

        fila = outcome.fila
        assert fila.acta == "N/A"
        assert fila.fecha == "N/A"
        assert "no se encontró el número del acta" in fila.observaciones

    def test_el_informe_trae_la_fila_para_juntarla_con_las_de_la_carpeta(self):
        informe = InventariarActa(lector_de_folios=Folios()).execute(
            acta_de_22_paginas(), document_name="a.pdf"
        ).as_dict()

        assert informe["folio_inicial"] == 1
        assert informe["folio_final"] == 21
        assert informe["cantidad_folios"] == 21
        assert informe["inventario_actas"][0]["acta"] == "ACTA ORDINARIA No. 022"
        # Y lo que lee la pantalla de Archivo.
        assert informe["asunto"] == "ACTA ORDINARIA No. 022 - Lunes 17 de diciembre de 2012"
        assert informe["fecha_final"] == "17-12-2012"
        json.dumps(informe)  # viaja entre procesos y por la API


# -----------------------------------------------------------------------------
#  Los adaptadores, sin red
# -----------------------------------------------------------------------------


class _Respuesta(io.BytesIO):
    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False


def _respuesta_de_chat(texto, entrada=7749, salida=209):
    return _Respuesta(
        json.dumps(
            {
                "choices": [{"message": {"content": texto}}],
                "usage": {"prompt_tokens": entrada, "completion_tokens": salida},
            }
        ).encode()
    )


def _http_error(codigo):
    return urllib.error.HTTPError("https://api.mistral.ai", codigo, "x", {}, io.BytesIO(b"{}"))


class TestElResumidorDeMistral:
    def test_manda_texto_al_chat_y_cuenta_los_tokens(self, monkeypatch):
        enviados = []

        def urlopen(request, timeout):
            enviados.append(json.loads(request.data))
            return _respuesta_de_chat("  Sesionó el consejo.  ")

        monkeypatch.setattr("urllib.request.urlopen", urlopen)
        consumo = Consumo()

        resumen = MistralResumidor("llave", consumo=consumo).resumir("texto del acta")

        assert resumen == "Sesionó el consejo."
        cuerpo = enviados[0]
        assert cuerpo["model"] == "mistral-small-latest"
        assert cuerpo["messages"][0]["content"] == INSTRUCCIONES
        assert cuerpo["messages"][1]["content"].endswith("texto del acta")
        assert "image" not in json.dumps(cuerpo)
        cuenta = consumo.as_dict()["proveedores"][PROVEEDOR]
        assert cuenta["tokens_entrada"] == 7749
        assert cuenta["tokens_salida"] == 209
        assert cuenta["paginas_facturadas"] == 0

    def test_las_instrucciones_prohiben_el_llamado_a_lista(self):
        assert "llamado a lista" in INSTRUCCIONES
        assert "quórum" in INSTRUCCIONES

    def test_un_429_se_reintenta_y_se_cuenta(self, monkeypatch):
        respuestas = [_http_error(429), _respuesta_de_chat("Listo.")]

        def urlopen(request, timeout):
            siguiente = respuestas.pop(0)
            if isinstance(siguiente, Exception):
                raise siguiente
            return siguiente

        monkeypatch.setattr("urllib.request.urlopen", urlopen)
        esperas = []
        consumo = Consumo()

        resumen = MistralResumidor("llave", consumo=consumo, esperar=esperas.append).resumir("t")

        assert resumen == "Listo."
        assert esperas == [2.0]
        assert consumo.as_dict()["proveedores"][PROVEEDOR]["rechazos"] == 1

    def test_tras_agotar_los_reintentos_se_rinde_sin_tumbar_nada(self, monkeypatch):
        def urlopen(request, timeout):
            raise _http_error(503)

        monkeypatch.setattr("urllib.request.urlopen", urlopen)
        resumidor = MistralResumidor(
            "llave", config=MistralResumenConfig(reintentos=1), esperar=lambda _: None
        )

        with pytest.raises(ResumenNoDisponible, match="HTTP 503"):
            resumidor.resumir("t")

    def test_una_llave_rechazada_cierra_el_resumidor_para_el_resto_del_trabajo(self, monkeypatch):
        llamadas = []

        def urlopen(request, timeout):
            llamadas.append(1)
            raise _http_error(401)

        monkeypatch.setattr("urllib.request.urlopen", urlopen)
        resumidor = MistralResumidor("llave mala")

        with pytest.raises(ResumenNoDisponible, match="rechazó la llave"):
            resumidor.resumir("acta 1")
        with pytest.raises(ResumenNoDisponible, match="rechazó la llave"):
            resumidor.resumir("acta 2")
        assert len(llamadas) == 1

    def test_una_caida_de_red_es_un_acta_sin_resumen(self, monkeypatch):
        def urlopen(request, timeout):
            raise urllib.error.URLError("sin conexión")

        monkeypatch.setattr("urllib.request.urlopen", urlopen)

        with pytest.raises(ResumenNoDisponible, match="no respondió"):
            MistralResumidor("llave").resumir("t")


class TestElLectorDeFolios:
    def test_apila_las_esquinas_en_una_sola_imagen(self):
        from PIL import Image

        recibidas = []

        def leer(imagen):
            recibidas.append(Image.open(io.BytesIO(imagen)).size)
            return "1\n\n21"

        texto = LectorDeFolios(leer).leer_esquinas(acta_de_22_paginas(), [2, 22])

        assert texto == "1\n\n21"
        # Una sola lectura -- una página facturada -- con las dos esquinas.
        assert len(recibidas) == 1
        ancho, alto = recibidas[0]
        assert ancho == 40 and alto > 2 * 20


# -----------------------------------------------------------------------------
#  Por dónde entra
# -----------------------------------------------------------------------------


class TestElTipoDeclarado:
    def test_se_puede_declarar(self):
        assert TipoPedido.parse("acta-consejo") is TipoPedido.ACTA_CONSEJO
        assert {"id": "acta-consejo", "label": "Actas de consejo"} in tipos_declarables()

    def test_el_despachador_lo_manda_a_su_habilidad_con_cualquier_accion(self, monkeypatch):
        from resolutions.api import worker

        llamadas = []
        monkeypatch.setattr(worker, "_acta_consejo_job", lambda payload: llamadas.append(payload) or {})
        monkeypatch.setattr(worker, "_anunciar", lambda *a, **k: None)

        for accion in ("inventory_file", "split", "inventory"):
            worker._despachar({"task": accion, "tipo": "acta-consejo", "settings": {}})

        assert len(llamadas) == 3


# -----------------------------------------------------------------------------
#  El Excel: una fila por PDF
# -----------------------------------------------------------------------------


def _fila(acta, archivo="1N°22.pdf"):
    return asdict(
        FilaDeActa(
            archivo=archivo,
            acta=acta,
            fecha="17-12-2012",
            folio_inicial=1,
            folio_final=21,
            folios=21,
            resumen="Sesionó el Consejo de Facultad.",
            carpeta="UPD3859771",
            pagina_inicial=1,
            pagina_final=21,
        )
    )


def _fila_de(hoja, numero):
    """Una fila del Excel por el nombre de sus columnas, no por su posición."""
    from resolutions.adapters.excel_actas import COLUMNAS

    return {
        titulo: hoja.cell(row=numero, column=columna).value
        for columna, (titulo, _clave, _ancho) in enumerate(COLUMNAS, start=1)
    }


def _indicador(hoja, nombre):
    """El valor de un indicador de la cabecera, por su nombre."""
    from resolutions.adapters.excel_actas import _INDICADORES, FILA_DE_INDICADORES

    columna = next(desde for titulo, desde, _hasta in _INDICADORES if titulo == nombre)
    return hoja.cell(row=FILA_DE_INDICADORES + 1, column=columna).value


class TestElExcel:
    def test_titulo_encabezados_y_filas_numeradas(self, tmp_path):
        from openpyxl import load_workbook

        from resolutions.adapters.excel_actas import (
            FILA_DE_ENCABEZADOS,
            FILA_DE_INDICADORES,
            escribir_inventario_de_actas,
        )

        destino = escribir_inventario_de_actas(
            [_fila("ACTA ORDINARIA No. 022"), _fila("ACTA ORDINARIA No. 023", "2N°23.pdf")],
            tmp_path / "actas__inventario.xlsx",
            carpeta="D:/ACTAS 2012",
        )

        hoja = load_workbook(destino).active
        assert hoja["A1"].value == "INVENTARIO DE ACTAS DE CONSEJO"
        assert "ACTAS 2012" in hoja["A2"].value
        # Los indicadores: actas, folios, periodo, por revisar y la carpeta.
        valores = FILA_DE_INDICADORES + 1
        assert _indicador(hoja, "ACTAS") == 2
        assert _indicador(hoja, "FOLIOS") == 42
        assert _indicador(hoja, "PERIODO") == "17-12-2012"
        assert _indicador(hoja, "POR REVISAR") == 0
        assert _indicador(hoja, "CARPETA") == "D:/ACTAS 2012"
        assert valores > 0

        encabezados = FILA_DE_ENCABEZADOS
        assert [celda.value for celda in hoja[encabezados]] == [
            "N°", "Carpeta", "Archivo PDF", "Resultado", "Páginas", "Acta", "Fecha de la sesión",
            "Folio inicial", "Folio final", "N° de folios", "Resumen del documento",
            "Piezas de la sesión", "Observaciones",
        ]
        primera = _fila_de(hoja, encabezados + 1)
        assert primera["N°"] == 1
        assert primera["Archivo PDF"] == "1N°22.pdf"
        assert primera["Páginas"] == "1–21"
        assert primera["Acta"] == "ACTA ORDINARIA No. 022"
        assert primera["Fecha de la sesión"] == "17-12-2012"
        assert (primera["Folio inicial"], primera["Folio final"], primera["N° de folios"]) == (
            1, 21, 21,
        )
        segunda = _fila_de(hoja, encabezados + 2)
        assert (segunda["N°"], segunda["Archivo PDF"]) == (2, "2N°23.pdf")
        assert "InventarioDeActas" in hoja.tables
        assert hoja.freeze_panes == f"A{encabezados + 1}"

    def test_el_periodo_va_de_la_mas_antigua_a_la_mas_reciente(self, tmp_path):
        from openpyxl import load_workbook

        from resolutions.adapters.excel_actas import (
            FILA_DE_INDICADORES,
            escribir_inventario_de_actas,
        )

        otra = {**_fila("ACTA No. 1"), "fecha": "28-01-2013"}
        sin_fecha = {**_fila("ACTA No. 2"), "fecha": "N/A", "observaciones": "sin fecha"}
        destino = escribir_inventario_de_actas(
            [otra, _fila("ACTA No. 3"), sin_fecha], tmp_path / "a__inventario.xlsx"
        )

        hoja = load_workbook(destino).active
        assert FILA_DE_INDICADORES > 0
        assert _indicador(hoja, "PERIODO") == "17-12-2012  a  28-01-2013"
        # Una con observaciones: una por revisar.
        assert _indicador(hoja, "POR REVISAR") == 1


# -----------------------------------------------------------------------------
#  Una carpeta de actas, un solo Excel
# -----------------------------------------------------------------------------


class TestElExcelDeLaCarpeta:
    @pytest.fixture
    def corrida(self, tmp_path):
        from dataclasses import replace

        from resolutions.api.folders import FolderRun, FolderRunner
        from resolutions.api.jobs import JobRegistry
        from resolutions.api.settings import Settings

        pytest.importorskip("openpyxl")
        settings = replace(Settings(), output_dir=tmp_path / "outputs")
        (tmp_path / "2012").mkdir()
        (tmp_path / "destino").mkdir()
        runner = FolderRunner(JobRegistry(), settings, run_job=None)
        run = FolderRun(id="r1", source=tmp_path / "2012", destination=tmp_path / "destino")
        return runner, run

    def test_cada_acta_entregada_entra_en_el_excel_de_la_carpeta(self, corrida):
        from openpyxl import load_workbook

        from resolutions.api.jobs import Job

        runner, run = corrida
        for numero, nombre in ((22, "1N°22.pdf"), (23, "2N°23.pdf")):
            job = Job(id=f"j{numero}", filename=nombre, source=run.source / nombre, owns_source=False)
            job.report = {"inventario_actas": [_fila(f"ACTA ORDINARIA No. 0{numero}", nombre)]}
            runner._deliver(run, job, run.source / nombre)

        from resolutions.adapters.excel_actas import FILA_DE_ENCABEZADOS

        # El destino es un espejo de la fuente: la carpeta elegida delante.
        assert Path(run.inventario) == run.destination / "2012" / "ACTAS 2012__inventario.xlsx"
        hoja = load_workbook(run.inventario).active
        primera = _fila_de(hoja, FILA_DE_ENCABEZADOS + 1)
        segunda = _fila_de(hoja, FILA_DE_ENCABEZADOS + 2)
        assert (primera["N°"], primera["Acta"]) == (1, "ACTA ORDINARIA No. 022")
        assert (segunda["N°"], segunda["Acta"]) == (2, "ACTA ORDINARIA No. 023")
        # La carpeta de cada fila es su ruta desde la raíz elegida.
        assert primera["Carpeta"] == "2012"
        # Sólo inventariar: el Excel en su carpeta espejo, sin ningún PDF.
        assert sorted(p.relative_to(run.destination).as_posix() for p in run.destination.rglob("*.*")) == [
            "2012/ACTAS 2012__inventario.xlsx"
        ]
        assert run.as_dict()["inventario"] == "ACTAS 2012__inventario.xlsx"
        assert run.as_dict()["inventario_filas"] == 2

    def test_se_reescribe_en_el_mismo_archivo(self, corrida):
        runner, run = corrida
        run.inventario_filas = [{**_fila("A"), "carpeta": "2012"}]
        runner._escribir_inventario_de_actas(run)
        primero = run.inventario

        run.inventario_filas.append({**_fila("B"), "carpeta": "2012"})
        runner._escribir_inventario_de_actas(run)

        assert run.inventario == primero
        assert len(list(run.destination.rglob("*.xlsx"))) == 1

    def test_reprocesar_la_carpeta_actualiza_su_excel_y_no_deja_un_dos(self, corrida):
        runner, run = corrida
        viejo = run.destination / "2012" / "ACTAS 2012__inventario.xlsx"
        viejo.parent.mkdir()
        viejo.write_bytes(b"de una corrida anterior")
        run.inventario_filas = [{**_fila("A"), "carpeta": "2012"}]

        runner._escribir_inventario_de_actas(run)

        assert Path(run.inventario) == viejo
        assert viejo.read_bytes() != b"de una corrida anterior"
        assert [p.name for p in run.destination.rglob("*.xlsx")] == ["ACTAS 2012__inventario.xlsx"]

    def test_cada_subcarpeta_lleva_su_excel_y_la_raiz_el_general(self, corrida):
        from openpyxl import load_workbook

        from resolutions.adapters.excel_actas import FILA_DE_ENCABEZADOS

        runner, run = corrida
        run.inventario_filas = [
            {**_fila("ACTA A", "a.pdf"), "carpeta": str(Path("2012") / "UPD1")},
            {**_fila("ACTA B", "b.pdf"), "carpeta": str(Path("2012") / "UPD2")},
        ]

        runner._escribir_inventario_de_actas(run)

        assert sorted(p.relative_to(run.destination).as_posix() for p in run.destination.rglob("*.xlsx")) == [
            "2012/ACTAS 2012__inventario.xlsx",
            "2012/UPD1/ACTAS UPD1__inventario.xlsx",
            "2012/UPD2/ACTAS UPD2__inventario.xlsx",
        ]
        suya = load_workbook(run.destination / "2012" / "UPD2" / "ACTAS UPD2__inventario.xlsx").active
        assert _fila_de(suya, FILA_DE_ENCABEZADOS + 1)["Acta"] == "ACTA B"

    def test_sin_actas_no_hay_excel(self, corrida):
        runner, run = corrida

        runner._escribir_inventario_de_actas(run)

        assert run.inventario is None


class TestLaDescargaDelExcelDeLaCarpeta:
    @staticmethod
    def _corrida(tmp_path, *, escrita=True):
        from resolutions.api import main
        from resolutions.api.folders import FolderRun, RunState

        run = FolderRun(id="carpeta1", source=tmp_path, destination=tmp_path, state=RunState.DONE)
        if escrita:
            hoja = tmp_path / "ACTAS__inventario.xlsx"
            hoja.write_bytes(b"planilla")
            run.inventario = str(hoja)
        main.contexto.folders._runs[run.id] = run
        return run

    def test_se_descarga(self, client, tmp_path):
        self._corrida(tmp_path)

        respuesta = client.get("/api/folder-runs/carpeta1/inventario.xlsx")

        assert respuesta.status_code == 200
        assert respuesta.content == b"planilla"
        assert "ACTAS__inventario.xlsx" in respuesta.headers["content-disposition"]

    def test_no_es_el_fuid_y_la_baja_tambien_el_tecnico(self, client, tecnico, tmp_path):
        self._corrida(tmp_path)

        respuesta = client.get("/api/folder-runs/carpeta1/inventario.xlsx", headers=tecnico)

        assert respuesta.status_code == 200

    def test_antes_de_escribirse_lo_dice(self, client, tmp_path):
        self._corrida(tmp_path, escrita=False)

        respuesta = client.get("/api/folder-runs/carpeta1/inventario.xlsx")

        assert respuesta.status_code == 404
        assert "todavía no tiene inventario" in respuesta.json()["detail"]

    def test_una_carpeta_que_no_existe(self, client):
        assert client.get("/api/folder-runs/nada/inventario.xlsx").status_code == 404

    def test_el_servicio_lo_anuncia(self, client):
        salud = client.get("/api/health").json()

        assert salud["api_revision"] >= 20
        assert {"acta-consejo", "folder-inventario"} <= set(salud["features"])
        assert any(tipo["id"] == "acta-consejo" for tipo in salud["document_types"])


# -----------------------------------------------------------------------------
#  «No estoy seguro»: el reconocedor también las distingue
# -----------------------------------------------------------------------------


class TestElReconocedor:
    def test_reconoce_un_acta_por_la_citacion_y_el_orden_del_dia(self):
        from resolutions.domain.doctype import DocumentType, classify_pages

        veredicto = classify_pages(enumerate(ACTA, start=1))

        assert veredicto.document_type is DocumentType.ACTA_CONSEJO

    def test_un_acta_de_la_correspondencia_no_es_un_acta_de_consejo(self):
        from resolutions.domain.doctype import DocumentType, classify_pages

        veredicto = classify_pages(
            [(1, "ACTA DE IRREGULARIDAD\nNIC 5825181\nSe encontró el medidor con el sello roto.")]
        )

        assert veredicto.document_type is not DocumentType.ACTA_CONSEJO

    def test_el_orquestador_cree_el_membrete_corrido_de_un_libro_antiguo(self):
        # Libro de 2000-2001 (UPD3859778): la muestra cae en hojas de cuerpo,
        # que sólo llevan «Consejo de Facultad / Acta No. 08-27-11-01» arriba.
        from resolutions.domain.validacion import validar_acta_consejo

        hoja = "Consejo de Facultad \nActa No. 08-27-11-01\n2\nde actualizar las cargas académicas"

        assert validar_acta_consejo([(8, hoja)]).valida

    def test_una_cita_a_mitad_de_hoja_no_es_un_membrete(self):
        from resolutions.domain.validacion import validar_acta_consejo

        hoja = "Texto del acta.\nSegún consta en el\nActa No. 021 se aprobó el calendario.\nCONSEJO DE FACULTAD"

        assert not validar_acta_consejo([(1, hoja)]).valida

    def test_el_despachador_las_inventaria_en_vez_de_partirlas(self, monkeypatch):
        from resolutions.api import worker
        from resolutions.domain.doctype import DocumentType

        llamadas = []
        monkeypatch.setattr(worker, "_recognise", lambda payload: DocumentType.ACTA_CONSEJO)
        monkeypatch.setattr(worker, "_acta_consejo_job", lambda payload: llamadas.append(1) or {})
        monkeypatch.setattr(worker, "_segment_job", lambda *a, **k: pytest.fail("no se parte"))
        monkeypatch.setattr(worker, "_split_job", lambda *a, **k: pytest.fail("no se parte"))
        monkeypatch.setattr(worker, "_anunciar", lambda *a, **k: None)

        worker._despachar({"task": "split", "tipo": "auto", "settings": {}})

        assert llamadas == [1]


# -----------------------------------------------------------------------------
#  Los libros: muchas actas en un PDF
# -----------------------------------------------------------------------------


def _hoja_de_libro(texto):
    return f"UNIVERSIDAD DE CARTAGENA\nFACULTAD DE ODONTOLOGIA\n{texto}\nCampus de Zaragocilla"


#: Un libro de 2012 como el de UPD3859772: cada acta con su portada separadora,
#: su citación y su encabezado; la segunda trae anexada un acta de un comité.
LIBRO = [
    "No. 002 -02 DE FEBRERO DE 2012",  # 1 portada
    _hoja_de_libro("CONSEJO DE FACULTAD\nSESION ORDINARIA\nFECHA: JUEVES 02 DE FEBRERO DE 2012\nORDEN DEL DIA:"),  # 2
    _hoja_de_libro("ACTA ORDINARIA No. 002\nJueves 2 de febrero de 2012\nCONSEJO DE FACULTAD"),  # 3
    _hoja_de_libro("Acta No. 002-2012\nConsejo de Facultad\nSe aprobó el calendario académico."),  # 4
    "No. 021 -14 DICIEMBRE DE 2012",  # 5 portada (el OCR se comió el «DE»)
    _hoja_de_libro("CONSEJO DE FACULTAD\nSESION EXTRAORDINARIA\nFECHA: VIERNES 14 DE DICIEMBRE DE 2012"),  # 6
    _hoja_de_libro("ACTA EXTRAORDINARIA No. 021\nViernes 14 de diciembre de 2012\nCONSEJO DE FACULTAD"),  # 7
    _hoja_de_libro("ACTA No 004\nCOMITE DE BIENESTAR\nFECHA: 3 de diciembre de 2012\nLUGAR: Decanatura"),  # 8 anexo
    _hoja_de_libro("Acta No. 021-2012\nSe aprobaron los microcurrículos."),  # 9
]


class TestLosLibros:
    def test_un_acta_por_pdf_sigue_siendo_el_pdf_entero(self):
        reparto = actas_del_pdf(ACTA)

        assert not reparto.es_libro
        assert reparto.tramos == (TramoDeActa(1, len(ACTA)),)

    def test_el_libro_se_reparte_por_sus_portadas_y_encabezados(self):
        reparto = actas_del_pdf(LIBRO)

        assert reparto.tramos == (TramoDeActa(1, 4), TramoDeActa(5, 9))
        assert reparto.previas is None

    def test_el_acta_de_un_comite_es_un_anexo_y_no_un_acta_del_libro(self):
        reparto = actas_del_pdf(LIBRO)

        assert all(not (t.desde <= 8 < t.hasta and t.desde == 8) for t in reparto.tramos)
        assert len(reparto.tramos) == 2

    def test_las_paginas_antes_de_la_primera_acta_se_apartan(self):
        reparto = actas_del_pdf(["CARTA DE REMISION DEL LIBRO", "Resolución del Ministerio", *LIBRO])

        assert reparto.previas == (1, 2)
        assert reparto.tramos[0] == TramoDeActa(3, 6)

    def test_los_formatos_de_2005_a_2009(self):
        libro = [
            _hoja_de_libro("ACTA CONSEJO DE FACULTAD ORDINARIO\nNo. 007\nFECHA:\nHORA:\nLUGAR:\n10 de agosto de 2005"),
            _hoja_de_libro("Pagina 02\nSe trató el presupuesto."),
            _hoja_de_libro("ACTA CONSEJO DE EACULTAD EXTRORDINARIO No.Oll\nFECHA: 19 DE ENERO DE 2009\nHORA: 9:00"),
        ]

        reparto = actas_del_pdf(libro)

        assert reparto.tramos == (TramoDeActa(1, 2), TramoDeActa(3, 3))
        primera, segunda = (leer_acta(libro[t.desde - 1 : t.hasta]) for t in reparto.tramos)
        assert (primera.numero, primera.sesion, primera.fecha) == ("007", "ORDINARIA", date(2005, 8, 10))
        assert (segunda.numero, segunda.sesion, segunda.fecha) == (
            "011", "EXTRAORDINARIA", date(2009, 1, 19),
        )

    def test_un_numero_de_pagina_pegado_no_abre_un_acta(self):
        # «Acta No. 0415»: el 04 con la página 15 pegada.
        libro = [_hoja_de_libro("ACTA No. 04\nFECHA: 15 de junio de 2001\nLUGAR: Decanatura"),
                 _hoja_de_libro("Acta No. 0415\nFECHA: x\nLUGAR: y")]

        assert len(actas_del_pdf(libro).tramos) == 1


class FolioPorActa:
    """Un lector de esquinas que contesta lo de cada acta, en orden."""

    def __init__(self, *respuestas):
        self.respuestas = list(respuestas)
        self.pedidas: list[list[int]] = []

    def leer_esquinas(self, source, paginas):
        self.pedidas.append(list(paginas))
        return self.respuestas.pop(0)


class TestInventariarUnLibro:
    def test_una_fila_por_acta_con_sus_paginas_folios_y_resumen(self):
        # Las portadas (hojas 1 y 5) separan sesiones y no llevan folio: no
        # cuentan como hojas foliadas aunque una mota en su esquina lo parezca.
        folios = FolioPorActa("1\n\n3", "4\n\n7")
        resumidor = Resumidor()
        outcome = InventariarActa(
            resumidor=resumidor, lector_de_folios=folios
        ).execute(Fuente(LIBRO, foliadas=range(1, 10)), document_name="libro.pdf", carpeta="UPD3859772")

        assert [f.acta for f in outcome.filas] == ["ACTA ORDINARIA No. 002", "ACTA EXTRAORDINARIA No. 021"]
        assert [(f.pagina_inicial, f.pagina_final) for f in outcome.filas] == [(1, 4), (5, 9)]
        assert [(f.folio_inicial, f.folio_final, f.folios) for f in outcome.filas] == [(1, 3, 3), (4, 7, 4)]
        assert {f.carpeta for f in outcome.filas} == {"UPD3859772"}
        assert folios.pedidas == [[2, 4], [6, 9]]
        assert all(f.resumen != "N/A" for f in outcome.filas)
        informe = outcome.as_dict()
        assert informe["records"] == 2
        assert informe["asunto"].startswith("LIBRO DE 2 ACTAS")

    def test_un_acta_mal_leida_del_libro_se_deduce_de_sus_vecinas(self):
        libro = LIBRO + [
            "No. 022 -17 DE DICIEMBRE DE 2012",
            _hoja_de_libro("ACTA ORDINARIA No. 022\nLunes 17 de diciembre de 2012\nCONSEJO DE FACULTAD"),
        ]
        # Sin las portadas (1, 5 y 10): tres, cuatro y una hoja foliadas.
        folios = FolioPorActa("1\n\n3", "381\n\n421", "8")
        outcome = InventariarActa(lector_de_folios=folios).execute(
            Fuente(libro, foliadas=range(1, 12)), document_name="libro.pdf"
        )

        medio = outcome.filas[1]
        assert (medio.folio_inicial, medio.folio_final) == (4, 7)
        assert "folios deducidos" in medio.observaciones

    def test_la_continuidad_no_cruza_de_una_carpeta_a_otra(self):
        from resolutions.application.acta_consejo import inventario_de_la_carpeta

        def fila(carpeta, archivo, inicial, final, folios):
            return {**_fila("ACTA", archivo), "carpeta": carpeta, "folio_inicial": inicial,
                    "folio_final": final, "folios": folios, "observaciones": ""}

        filas = inventario_de_la_carpeta(
            [
                fila("UPD1", "1.pdf", 1, 20, 20),
                fila("UPD2", "1.pdf", "N/A", "N/A", 10),
                fila("UPD2", "2.pdf", 31, 40, 10),
            ]
        )

        assert filas[1]["carpeta"] == "UPD2" and filas[1]["folio_inicial"] == "N/A"


# -----------------------------------------------------------------------------
#  Lo que ve la consola mientras se lee
# -----------------------------------------------------------------------------


class _Eventos:
    def __init__(self):
        self.recibidos = []

    def emit(self, evento):
        self.recibidos.append(evento)


class _OcrRoto:
    def read(self, imagen):
        raise RuntimeError("tesseract no respondió")


class TestLasPaginasLeidasNoVanARevision:
    """Todas las páginas de todas las actas salían «enviadas a revisión».

    El aviso de cada página leída no decía de dónde salió su texto, y la cinta
    de progreso toma una página sin procedencia por ilegible.
    """

    @staticmethod
    def _cinta(eventos):
        from dataclasses import asdict

        from resolutions.api.jobs import JobProgress

        progreso = JobProgress()
        for evento in eventos.recibidos:
            progreso.apply({k: v for k, v in asdict(evento).items() if v is not None})
        return progreso

    def test_ninguna_pagina_leida_se_marca(self):
        eventos = _Eventos()
        InventariarActa(ocr=OcrDeCasa(), progress=eventos).execute(
            acta_de_22_paginas(), document_name="a.pdf"
        )

        progreso = self._cinta(eventos)
        assert "x" not in "".join(progreso.marks)
        assert progreso.failed_pages == 0
        # La portada pasó por el OCR y las demás salieron de su capa.
        assert progreso.marks[0] == "f"
        assert set(progreso.marks[1:]) == {"t"}

    def test_solo_la_que_no_se_pudo_leer_va_a_revision_y_dice_por_que(self):
        eventos = _Eventos()
        InventariarActa(ocr=_OcrRoto(), progress=eventos).execute(
            acta_de_22_paginas(), document_name="a.pdf"
        )

        progreso = self._cinta(eventos)
        assert progreso.marks.count("x") == 1
        assert progreso.review == {1: "falló el OCR de la página"}
