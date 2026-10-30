"""Todo lo de la Universidad de Cartagena lleva el mismo FUID (2026-10-30).

El operador: «todo lo que sea de Universidad de Cartagena lleva ese inventario,
estandarízalo». La plantilla es la suya (``assets/fuid-fo-gd-008.xlsx``), con la
oficina de cada carpeta en la cabecera, y ningún PDF se queda sin su fila: el
que su habilidad no sabe partir se inventaría entero, con el asunto y los años
de su nombre, como el inventario del propio cliente.
"""

from __future__ import annotations

import time

import pytest

pytest.importorskip("openpyxl")

from openpyxl import load_workbook  # noqa: E402

from resolutions.adapters.fuid_de_la_entrega import (  # noqa: E402
    FILA_HECHA,
    PDF_ENTERO,
    anios_del_nombre,
    asunto_del_nombre,
    fila_del_fuid,
)
from resolutions.application.inventory_document import template_for  # noqa: E402
from resolutions.domain.doctype import DocumentType  # noqa: E402


@pytest.mark.parametrize(
    ("archivo", "anios"),
    [
        ("UPD3860000-EXAMENES PREPARATORIOS DE GRADO 1924-46.pdf", (1924, 1946)),
        ("UPD 3859599 HISTORIAS ACADEMICAS 19441969.pdf", (1944, 1969)),
        ("1HISTORIAS ACADEMICAS 1983-1977.pdf", (1977, 1983)),
        ("UPD3859739-1HISTORIAS ACADEMICAS 2002-2.pdf", (2002, 2002)),
        ("UPD3859806-ACTAS DE GRADO POSTGRADO-PREGRADO 1997-1998-N°1 A.pdf", (1997, 1998)),
        ("UPD3859729-1HISTORIAS ACADEMICAS 14 DE DICIEMBRE 1988.pdf", (1988, 1988)),
        ("1998-02 ACTAS.pdf", (1998, 2002)),
        # Un número que no es un año no se adivina: ese extremo queda sin fecha.
        ("UPD3859999-EXAMENES PREPARATORIOS DE GRADO 1926-1395.pdf", (1926, None)),
        ("UPD0077290- HISTORIAS ACADEMICAS (1190-1995).pdf", (None, 1995)),
        ("ACTAS SIN FECHA.pdf", (None, None)),
    ],
)
def test_los_anios_del_nombre(archivo, anios):
    assert anios_del_nombre(archivo) == anios


@pytest.mark.parametrize(
    ("archivo", "asunto"),
    [
        ("UPD3860000-EXAMENES PREPARATORIOS DE GRADO 1924-46.pdf", "EXAMENES PREPARATORIOS DE GRADO 1924-46"),
        ("UPD 3859601 HISTORIAS ACADEMICAS 1980-1992.pdf", "HISTORIAS ACADEMICAS 1980-1992"),
        ("UPD0077188 - MATRICULAS ESPECIALES 5 Y 6 AÑO 1991 - 1992.pdf", "MATRICULAS ESPECIALES 5 Y 6 AÑO 1991 - 1992"),
        ("UPD3859752-1HISTORIAS ACADEMICAS RETIRADOS 1989.pdf", "HISTORIAS ACADEMICAS RETIRADOS 1989"),
        ("1HISTORIAS ACADEMICAS 1946-1990.pdf", "HISTORIAS ACADEMICAS 1946-1990"),
    ],
)
def test_el_asunto_del_nombre_sin_la_upd_ni_el_orden(archivo, asunto):
    assert asunto_del_nombre(archivo) == asunto


def test_la_fila_de_un_pdf_entero():
    fila = fila_del_fuid(
        {
            "tipo": PDF_ENTERO,
            "archivo": "UPD 3859597 HISTORIAS ACADEMICAS 1948-1960.pdf",
            "carpeta": "17. FACULTAD DE MEDICINA 3261- 3268 P\\CAJA00C003261",
            "folios": 187,
            "pagina_inicial": 1,
            "pagina_final": 187,
            "entero": True,
            "observaciones": "No se reconoció la ficha de ningún estudiante",
        }
    )
    assert fila.asunto == "HISTORIAS ACADEMICAS 1948-1960"
    assert (fila.fecha_inicial, fila.fecha_final) == ("01/01/1948", "31/12/1960")
    assert (fila.caja, fila.carpeta) == ("00C003261", "UPD3859597")
    assert (fila.folios, fila.folios_siar) == (187, 187)
    assert fila.notas == (
        "No se reconoció la ficha de ningún estudiante; Archivo: UPD 3859597 HISTORIAS ACADEMICAS 1948-1960.pdf"
    )


def _archivo_ilegible(**cambios):
    """La fila que deja «Archivos ya separados» de un manuscrito que no se deja leer."""
    datos = {
        "orden": 1,
        "asunto": "Documento sin identificar",
        "fecha_inicial": "N/A",
        "fecha_final": "N/A",
        "folios": 199,
        "folios_siar": 199,
        "notas": "",
        "pagina_inicial": 1,
    }
    return {**datos, **cambios}


def test_el_manuscrito_que_no_se_lee_toma_asunto_y_anios_de_su_nombre():
    fila = fila_del_fuid(
        {
            "tipo": FILA_HECHA,
            "fuid": _archivo_ilegible(),
            "carpeta": "20.ACTAS DE GRADO 3362",
            "archivo": "UPD3860000-EXAMENES PREPARATORIOS DE GRADO 1924-46.pdf",
            "pagina_inicial": 1,
            "entero": True,
        }
    )
    assert fila.asunto == "EXAMENES PREPARATORIOS DE GRADO 1924-46"
    assert (fila.fecha_inicial, fila.fecha_final) == ("01/01/1924", "31/12/1946")
    assert fila.carpeta == "UPD3860000"
    assert "del nombre del archivo" in fila.notas
    assert fila.notas.endswith("Archivo: UPD3860000-EXAMENES PREPARATORIOS DE GRADO 1924-46.pdf")


def test_lo_que_si_se_leyo_del_papel_no_se_cambia_por_el_nombre():
    fila = fila_del_fuid(
        {
            "tipo": FILA_HECHA,
            "fuid": _archivo_ilegible(asunto="ACTA DE GRADO", fecha_inicial="04-07-1924", fecha_final="22-05-1946"),
            "carpeta": "20.ACTAS DE GRADO 3362",
            "archivo": "UPD3860000-EXAMENES PREPARATORIOS DE GRADO 1924-46.pdf",
            "entero": True,
        }
    )
    assert (fila.asunto, fila.fecha_inicial, fila.fecha_final) == ("ACTA DE GRADO", "04-07-1924", "22-05-1946")
    assert "del nombre" not in fila.notas


def test_una_resolucion_partida_no_toma_nada_del_nombre():
    """Sólo el archivo entero: una resolución sin fecha legible se queda N/A."""
    fila = fila_del_fuid(
        {
            "tipo": FILA_HECHA,
            "fuid": _archivo_ilegible(asunto="RESOLUCION 00031", notas="Archivo: 02_RESOLUCION-00031.pdf"),
            "carpeta": "00C003303",
            "archivo": "UPD3859990-RESOLUCIONES INTERNAS 2010 N°30-31-32.pdf",
        }
    )
    assert (fila.fecha_inicial, fila.fecha_final) == ("N/A", "N/A")


# -- una sola plantilla -------------------------------------------------------


@pytest.mark.parametrize(
    "tipo",
    [DocumentType.ACTA, DocumentType.RESOLUCION, DocumentType.MATRICULA, DocumentType.DESCONOCIDO, DocumentType.DIPLOMA],
)
def test_todas_las_habilidades_de_la_universidad_usan_la_plantilla_del_operador(tipo):
    """La misma hoja, el mismo bloque de firmas debajo de las filas y la misma fila fija."""
    hoja = load_workbook(template_for(tipo)).active
    assert hoja["C3"].value == "FORMATO ÚNICO DE INVENTARIO DOCUMENTAL - FUID"
    assert hoja["P1"].value == "CÓDIGO: FO-GD-008"
    assert hoja.freeze_panes == "A14"
    assert hoja["A43"].value == "Elaborado por:"


def test_los_diplomas_conservan_la_cabecera_de_la_secretaria_general():
    assert load_workbook(template_for(DocumentType.DIPLOMA)).active["B8"].value == (
        "OFICINA PRODUCTORA: SECRETARÍA GENERAL - RECTORÍA"
    )


# -- el trabajo de historias sin fichas -----------------------------------------


def test_historias_sin_ninguna_ficha_se_inventarian_enteras(tmp_path):
    """Un formato que la habilidad no conoce no se pierde: una fila y el PDF entero."""
    pymupdf = pytest.importorskip("pymupdf")
    from resolutions.api.settings import Settings
    from resolutions.api.worker.historias import historias_academicas_job

    pdf = tmp_path / "UPD 3859601 HISTORIAS ACADEMICAS 1980-1992.pdf"
    with pymupdf.open() as documento:
        for _ in range(3):
            # Hojas de notas: no abren expediente y no hace falta releerlas.
            documento.new_page().insert_text((72, 72), "REGISTRO DE ESTUDIOS ASIGNATURAS PROMEDIO 1980")
        documento.save(pdf)
    ajustes = Settings(output_dir=tmp_path / "salida").as_worker_payload()
    for clave in ("anthropic_api_key", "gemini_api_key", "mistral_api_key"):
        ajustes[clave] = None

    informe = historias_academicas_job(
        {
            "job_id": "historias-sin-fichas",
            "source": str(pdf),
            "filename": pdf.name,
            "operator": "prueba",
            "task": "both",
            "lectura": "local",
            "tipo": "historia-academica",
            "oficina": "FACULTAD DE MEDICINA",
            "settings": ajustes,
        }
    )

    assert informe["historias"] == 0
    # Sin PDF partidos: la corrida de carpeta entrega el original con su nombre.
    assert not informe.get("outputs")
    (fila,) = informe["inventario_actas"]
    assert (fila["tipo"], fila["folios"], fila["entero"]) == (PDF_ENTERO, 3, True)
    convertida = fila_del_fuid(fila)
    assert convertida.asunto == "HISTORIAS ACADEMICAS 1980-1992"
    assert (convertida.fecha_inicial, convertida.fecha_final) == ("01/01/1980", "31/12/1992")
    assert informe["inventario"]


# -- la pantalla ------------------------------------------------------------------


def test_la_carpeta_lleva_la_oficina_y_la_plantilla_que_manda_la_pantalla(client, tmp_path):
    """La pantalla manda «tipoDocumento»: leyendo sólo «tipo_documento» se perdía."""
    origen = tmp_path / "17. FACULTAD DE MEDICINA"
    destino = tmp_path / "destino"
    origen.mkdir()
    destino.mkdir()

    respuesta = client.post(
        "/api/folder-runs",
        json={
            "source": str(origen),
            "destination": str(destino),
            "task": "both",
            "tipoDocumento": "kingspan",
            "oficina": "  facultad de   medicina ",
        },
    )
    assert respuesta.status_code == 201, respuesta.text
    corrida = respuesta.json()
    assert corrida["tipo_documento"] == "kingspan"
    assert corrida["oficina"] == "FACULTAD DE MEDICINA"

    client.post(f"/api/folder-runs/{corrida['id']}/stop")
    limite = time.monotonic() + 10
    while time.monotonic() < limite:
        if client.get(f"/api/folder-runs/{corrida['id']}").json()["state"] != "running":
            break
        time.sleep(0.05)


def test_el_reconocimiento_dice_que_oficina_llevara_el_fuid(client, tmp_path):
    carpeta = tmp_path / "14. ENFERMERIA 3277-3283 P"
    carpeta.mkdir()
    respuesta = client.post("/api/reconocer-carpeta", json={"ruta": str(carpeta)})
    assert respuesta.status_code == 200, respuesta.text
    assert respuesta.json()["oficina"] == "FACULTAD DE ENFERMERIA"


# -- los folios del inventario del cliente (opción B, 2026-10-30) --------------------


def test_los_folios_del_cliente_salen_de_su_excel(tmp_path):
    from openpyxl import Workbook

    from resolutions.adapters.inventario_del_cliente import folios_por_upd

    libro = Workbook()
    hoja = libro.active
    hoja.append(["ITEM", "NOMBRE O ASUNTOS", "FOLIOS ", "IMÁGENES", "CAJA", "UPD "])
    hoja.append([1, "EXAMENES PREPARATORIOS DE GRADO 1926-1395", 154, 305, "00C003362", "UPD3859999"])
    hoja.append([2, "EXAMENES PREPARATORIOS DE GRADO 1924-46", 101, 199, "00C003362", "UPD3860000"])
    libro.save(tmp_path / "ACTAS DE GRADO MEDICINA.xlsx")
    assert folios_por_upd(tmp_path / "ACTAS DE GRADO MEDICINA.xlsx") == {"UPD3859999": 154, "UPD3860000": 101}


def test_el_pdf_entero_lleva_las_hojas_del_cliente_y_las_paginas_en_siar():
    """UPD3860000: 101 hojas en el inventario del cliente y 199 páginas escaneadas."""
    fila = fila_del_fuid(
        {
            "tipo": FILA_HECHA,
            "fuid": _archivo_ilegible(),
            "carpeta": "20.ACTAS DE GRADO 3362",
            "archivo": "UPD3860000-EXAMENES PREPARATORIOS DE GRADO 1924-46.pdf",
            "entero": True,
            "folios_del_cliente": 101,
        }
    )
    assert (fila.folios, fila.folios_siar) == (101, 199)
    assert "las hojas del inventario del cliente" in fila.notas
