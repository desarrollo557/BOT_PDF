"""La fila del inventario oficial de Kingspan que le toca a una UPD.

Los textos y la secuencia de números están copiados de la UPD3869754, un paquete
real de notas internas de mayo de 2024 cuya portada, escrita a mano, dice
«14161 - 14368». Lo que se comprueba es lo que pidió el operador: el primer y el
último documento del paquete, cada uno con su fecha, en una sola fila.
"""

from __future__ import annotations

from datetime import date

import openpyxl

from resolutions.api.worker.kingspan import (
    HOJA,
    INVENTORY_SUFFIX,
    PRIMERA_FILA,
    celdas_kingspan,
    escritor,
    inventario_general,
)
from resolutions.domain.kingspan import (
    Documento,
    FilaKingspan,
    documento_de,
    filas_de,
    leer_numero,
)

#: La página 2 de la UPD3869754 tal como la trae su capa de texto.
NOTA_14368 = """NOTA INTERNA DE CONTABILIDAD
KINGSPAN PANELES AISLADOS S.A.S
Número:
OOl-NI-00014368
NIT 900447906-1
26/05/2024
PRINCIPAL
D.Cruce/M.
R44E-0029021
"""

#: La página 49: la fecha va detrás de otros datos del encabezado.
NOTA_14417 = """NOTA INTERNA DE CONTABILIDAD
KINGSPAN PANELES AISLADOS S.A.S
NIT 900447906-1
Ntímero:
OOl-NI-00014417
PRINCIPAL
Fecha:
26/05/2024
"""

#: La secuencia de notas internas de la UPD3869754 en el orden de sus páginas.
SECUENCIA = [
    14368, 14369, 14370, 14355, 14325, 14354, 14363, 14417, 14292, 14406, 14405,
    14379, 14373, 14372, 14371, 14359, 14358, 14357, 14349, 14347, 14345, 14343,
    14342, 14341, 14326, 14306, 14299, 14296, 14295, 14290, 14281, 14278, 14270,
    14260, 14257, 14249, 14246, 14244, 14237, 14234, 14233, 14232, 14217, 14207,
    14206, 14205, 14202, 14190, 14189, 14176, 14175, 14174, 14173, 14172, 14171,
    14169, 14168, 14167, 14166, 14165, 14162, 14161,
]


def _notas(numeros, fechas=None, serie="NI"):
    fechas = fechas or {}
    return [
        Documento(pagina=i + 2, serie=serie, numero=n, fecha=fechas.get(n, date(2024, 5, 26)))
        for i, n in enumerate(numeros)
    ]


class TestElNumeroImpreso:
    """La capa de texto de estos escaneos confunde letras y cifras."""

    def test_la_o_y_la_ele_del_consecutivo(self):
        assert leer_numero(NOTA_14368)[:2] == ("NI", 14368)

    def test_ceros_leidos_como_letra_o(self):
        assert leer_numero("Número:\n001-NI-OOO14189")[:2] == ("NI", 14189)

    def test_un_correlativo_partido_por_un_espacio(self):
        assert leer_numero("Niimero:\n001-NI-OOOl424 6")[:2] == ("NI", 14246)

    def test_un_espacio_no_pega_lo_que_viene_detras(self):
        assert leer_numero("Número:\nOOl-NI-00014417 1 n 1")[:2] == ("NI", 14417)

    def test_manda_el_numero_del_encabezado_y_no_una_cita_del_cuerpo(self):
        texto = "D.Cruce 001-CCF-00019746\nNúmero:\n001-NI-00014189\nFecha:\n1/05/2024"
        assert leer_numero(texto)[:2] == ("NI", 14189)

    def test_el_documento_soporte_lleva_otro_formato(self):
        assert leer_numero("Número DSS 1747\nFecha de Operación 5/05/2024")[:2] == ("DSS", 1747)

    def test_una_pagina_sin_numero(self):
        assert leer_numero("DIAN FACTURA ELECTRÓNICA DE VENTA") is None


class TestLaFechaDeCadaDocumento:
    def test_la_del_encabezado(self):
        documento = documento_de(NOTA_14368, 2)
        assert documento.fecha == date(2024, 5, 26)
        assert documento.codigo == "001-NI-00014368"

    def test_la_que_va_detras_del_numero(self):
        assert documento_de(NOTA_14417, 49).fecha == date(2024, 5, 26)

    def test_anio_mes_dia(self):
        documento = documento_de("Número DSS 1064\nFecha 2023/10/14", 3)
        assert documento.fecha == date(2023, 10, 14)
        assert documento.codigo == "DSS 1064"


class TestUnaFilaPorUpd:
    """El rango es de punta a punta del paquete, no por tramos consecutivos."""

    def test_la_upd3869754_da_lo_que_dice_su_portada(self):
        fechas = {14368: date(2024, 5, 26), 14161: date(2024, 5, 5)}
        filas = filas_de(_notas(SECUENCIA, fechas))
        assert len(filas) == 1
        fila = filas[0]
        assert fila.asunto == "NOTAS INTERNAS"
        assert (fila.desde, fila.hasta) == ("001-NI-00014161", "001-NI-00014368")

    def test_escaneada_al_reves_las_fechas_siguen_a_su_documento(self):
        # Fecha inicial la del 14161 y final la del 14368, aunque la primera
        # página sea la del 14368: cada punta conserva la suya.
        fechas = {14368: date(2024, 5, 26), 14161: date(2024, 5, 5)}
        fila = filas_de(_notas(SECUENCIA, fechas))[0]
        assert fila.fecha_inicial == date(2024, 5, 5)
        assert fila.fecha_final == date(2024, 5, 26)

    def test_las_fechas_no_son_la_minima_y_la_maxima(self):
        # El oficial tiene 65 filas con la final anterior a la inicial: la
        # UPD3859157 va del 21045 (21/03/2024) al 21111 (18/03/2024).
        fechas = {21045: date(2024, 3, 21), 21080: date(2024, 3, 1), 21111: date(2024, 3, 18)}
        fila = filas_de(_notas([21045, 21080, 21111], fechas, serie="CEE"))[0]
        assert fila.fecha_inicial == date(2024, 3, 21)
        assert fila.fecha_final == date(2024, 3, 18)
        assert fila.desde == "001-CEE-00021045"

    def test_lo_que_queda_fuera_del_rango_va_a_notas(self):
        fila = filas_de(_notas(SECUENCIA))[0]
        assert fila.notas.startswith("9 documento(s) fuera del rango 14161-14368")
        assert "14417" in fila.notas

    def test_los_anexos_no_cuentan_en_el_rango(self):
        documentos = _notas(SECUENCIA) + _notas([19746, 20368, 20287], serie="CCF")
        filas = filas_de(documentos)
        assert len(filas) == 1
        assert filas[0].desde == "001-NI-00014161"

    def test_dos_asuntos_parejos_son_dos_filas(self):
        # Como la UPD1923787 del oficial: notas bancarias y egresos.
        documentos = _notas(range(1300, 1310), serie="NB") + _notas(range(9900, 9910), serie="CEE")
        filas = filas_de(documentos)
        assert [fila.asunto for fila in filas] == ["NOTAS BANCARIAS", "COMPROBANTES DE EGRESO ELECTRONICO"]

    def test_un_numero_repetido_cuenta_una_vez(self):
        filas = filas_de(_notas([100, 101, 101, 102]))
        assert len(filas[0].documentos) == 3

    def test_sin_fecha_se_dice(self):
        documentos = [Documento(pagina=1, serie="NI", numero=5, fecha=None)]
        assert "no se leyó la fecha" in filas_de(documentos)[0].notas

    def test_sin_documentos_no_hay_filas(self):
        assert filas_de([]) == []


class TestCajaYUpd:
    """Del nombre con que llegó el trabajo, no del archivo que guardó el servicio."""

    def test_una_subida_suelta(self):
        from types import SimpleNamespace

        from resolutions.api.worker.kingspan import _caja, _upd

        taller = SimpleNamespace(
            name="UPD3869754.pdf",
            delivered_to=None,
            source="/data/uploads/C157F939EEE64BDF8D21A5249B58EC6C.pdf",
        )
        assert _upd(taller.name) == "UPD3869754"
        assert _caja(taller) == "N/A"

    def test_una_carpeta_trae_la_caja_en_la_ruta(self):
        from types import SimpleNamespace

        from resolutions.api.worker.kingspan import _caja

        taller = SimpleNamespace(
            name="100C000235/UPD3869754.pdf", delivered_to=None, source="/tmp/copia.pdf"
        )
        assert _caja(taller) == "100C000235"


class TestSeConfirmaEnLaImagen:
    """La capa de texto del escaneo puede cambiar una cifra: la p.107 de la
    UPD3869754 dice 14346 y la capa traía 14345. Lo que va al inventario se
    relee en la imagen."""

    def _taller(self, lecturas):
        from types import SimpleNamespace

        class Ocr:
            def read(self, imagen):
                return SimpleNamespace(text=lecturas.get(imagen, ""))

        control = SimpleNamespace(check=lambda: None)
        return SimpleNamespace(ocr=lambda: Ocr(), control=control)

    def _fuente(self):
        from types import SimpleNamespace

        return SimpleNamespace(render=lambda pagina, banda, dpi: pagina)

    def test_una_punta_mal_leida_se_corrige_y_se_anota(self):
        from resolutions.api.worker.kingspan import _confirmar

        documentos = _notas([14368, 14300, 14161])  # páginas 2, 3 y 4
        lecturas = {2: "Número:\n001-NI-00014388\nFecha:\n26/05/2024", 4: "Número:\n001-NI-00014161\nFecha:\n26/05/2024"}
        corregidos, notas, releidas = _confirmar(self._taller(lecturas), self._fuente(), documentos)

        assert releidas == 2
        assert corregidos[0].numero == 14388
        assert filas_de(corregidos)[0].hasta == "001-NI-00014388"
        assert notas["NI"] == [
            "verificar p.2: la imagen dice 001-NI-00014388 del 26/05/2024 y el texto del escaneo 001-NI-00014368"
        ]

    def test_si_la_imagen_no_se_lee_se_queda_la_capa(self):
        from resolutions.api.worker.kingspan import _confirmar

        documentos = _notas([14368, 14161])
        corregidos, notas, _ = _confirmar(self._taller({}), self._fuente(), documentos)
        assert corregidos == documentos
        assert notas == {}


class TestEntidadesYActa:
    """Las entidades van fijas y el acta la elige el operador (2026-10-26)."""

    def _celdas(self, **extra):
        return celdas_kingspan(
            FilaKingspan(asunto="NOTAS INTERNAS"),
            orden=1,
            caja="N/A",
            upd="N/A",
            settings=extra.pop("settings", {}),
            inventariado=date(2026, 10, 26),
            **extra,
        )

    def test_remitente_y_productora_siempre_kingspan(self):
        celdas = self._celdas(
            settings={"fuid_entidad_remitente": "OTRA", "fuid_entidad_productora": "OTRA"}
        )
        assert (celdas[2], celdas[3]) == ("KINGSPAN SAS", "KINGSPAN SAS")

    def test_el_acta_elegida_va_en_su_columna(self):
        assert self._celdas(acta="092")[25] == "092"

    def test_sin_acta_va_n_a(self):
        assert self._celdas()[25] == "N/A"

    def test_el_servicio_rechaza_un_acta_que_no_existe(self):
        import pytest
        from fastapi import HTTPException

        from resolutions.api.routers.trabajos import _acta

        assert _acta("092") == "092"
        assert _acta("") is None
        assert _acta(None) is None
        with pytest.raises(HTTPException) as error:
            _acta("999")
        assert error.value.status_code == 422

    def test_las_actas_son_las_del_oficial(self):
        from resolutions.domain.kingspan import ACTAS_DE_TRANSFERENCIA

        assert ACTAS_DE_TRANSFERENCIA == ("012", "015", "070", "088", "092", "094")


class TestLaPlanilla:
    """Se escribe sobre el inventario oficial, desde su fila 7."""

    def test_la_fila_cae_en_sus_columnas(self, tmp_path):
        fila = FilaKingspan(
            asunto="NOTAS INTERNAS",
            desde="001-NI-00014161",
            hasta="001-NI-00014368",
            fecha_inicial=date(2024, 5, 5),
            fecha_final=date(2024, 5, 26),
        )
        celdas = celdas_kingspan(
            fila,
            orden=1,
            caja="100C000235",
            upd="UPD3869754",
            settings={},
            inventariado=date(2026, 10, 26),
        )
        assert len(celdas) == 27
        destino = escritor().write_cells([celdas], tmp_path / "inventario.xlsx")

        hoja = openpyxl.load_workbook(destino)[HOJA]
        assert hoja["K6"].value == "DESDE"
        fila7 = {c.column_letter: c.value for c in hoja[PRIMERA_FILA]}
        assert fila7["C"] == "KINGSPAN SAS"
        assert fila7["J"] == "NOTAS INTERNAS"
        assert (fila7["K"], fila7["L"]) == ("001-NI-00014161", "001-NI-00014368")
        assert fila7["M"].date() == date(2024, 5, 5)
        assert (fila7["O"], fila7["P"]) == ("100C000235", "UPD3869754")
        assert fila7["W"] == "N/A"
        assert fila7["AA"] == "N/A"

    def test_cada_celda_lleva_el_formato_del_oficial(self, tmp_path):
        # Igual que la primera fila de «1. INVENTARIO DOCUMENTAL KINGSPAN.xlsx»:
        # fechas mm-dd-yy, texto en CAJA y ACTA, General en el resto.
        from resolutions.api.worker.kingspan import PLANTILLA

        celdas = celdas_kingspan(
            FilaKingspan(asunto="NOTAS INTERNAS", fecha_inicial=date(2024, 5, 5)),
            orden=1,
            caja="100C000235",
            upd="UPD3869754",
            settings={},
            inventariado=date(2026, 10, 26),
        )
        escrito = openpyxl.load_workbook(escritor().write_cells([celdas], tmp_path / "i.xlsx"))[HOJA]
        plantilla = openpyxl.load_workbook(PLANTILLA)[HOJA]
        assert [c.number_format for c in escrito[PRIMERA_FILA]] == [
            c.number_format for c in plantilla[PRIMERA_FILA]
        ]
        assert escrito["M7"].number_format == "mm-dd-yy"
        assert escrito["O7"].number_format == "@"

    def test_notas_va_como_en_el_oficial_aunque_haya_algo_que_revisar(self):
        celdas = celdas_kingspan(
            FilaKingspan(asunto="NOTAS INTERNAS", notas="9 documento(s) fuera del rango"),
            orden=1,
            caja="N/A",
            upd="N/A",
            settings={},
            inventariado=date(2026, 10, 26),
        )
        assert celdas[22] == "N/A"

    def test_la_fila_escrita_se_ve(self, tmp_path):
        # El oficial traía un filtro que ocultaba 1.436 filas, la 7 entre
        # ellas: el inventario llegaba con su fila escondida y parecía vacío.
        celdas = celdas_kingspan(
            FilaKingspan(asunto="NOTAS INTERNAS"),
            orden=1,
            caja="N/A",
            upd="N/A",
            settings={},
            inventariado=date(2026, 10, 26),
        )
        destino = escritor().write_cells([celdas, celdas], tmp_path / "inventario.xlsx")
        hoja = openpyxl.load_workbook(destino)[HOJA]
        assert not hoja.row_dimensions[7].hidden
        assert not hoja.row_dimensions[8].hidden
        assert hoja.auto_filter.ref is None

    def test_la_plantilla_no_trae_filas_ocultas_ni_filtro(self):
        from resolutions.api.worker.kingspan import PLANTILLA

        hoja = openpyxl.load_workbook(PLANTILLA)[HOJA]
        assert not [n for n, d in hoja.row_dimensions.items() if d.hidden]
        assert hoja.auto_filter.ref is None
        assert hoja.sheet_view.selection[0].activeCell == "A7"

    def test_el_general_de_la_carpeta_ordena_por_caja_y_reinicia_el_orden(self, tmp_path):
        def upd(caja, numero, desde):
            fila = FilaKingspan(asunto="NOTAS INTERNAS", desde=desde, hasta=desde)
            celdas = celdas_kingspan(
                fila,
                orden=1,
                caja=caja,
                upd=numero,
                settings={},
                inventariado=date(2026, 10, 26),
            )
            carpeta = tmp_path / caja / numero
            escritor().write_cells([celdas], carpeta / f"{numero}{INVENTORY_SUFFIX}")

        upd("100C000236", "UPD0000003", "C")
        upd("100C000235", "UPD0000002", "B")
        upd("100C000235", "UPD0000001", "A")

        general = inventario_general(tmp_path)
        assert general.name == "00.INVENTARIO DOCUMENTAL KINGSPAN.xlsx"
        hoja = openpyxl.load_workbook(general)[HOJA]
        filas = [
            (fila[0], fila[14], fila[15])
            for fila in hoja.iter_rows(min_row=PRIMERA_FILA, max_col=27, values_only=True)
        ]
        assert filas == [
            (1, "100C000235", "UPD0000001"),
            (2, "100C000235", "UPD0000002"),
            (1, "100C000236", "UPD0000003"),
        ]

    def test_sin_inventarios_no_hay_general(self, tmp_path):
        assert inventario_general(tmp_path) is None

    def test_los_ajustes_mandan_sobre_los_fijos(self):
        celdas = celdas_kingspan(
            FilaKingspan(asunto="X"),
            orden=1,
            caja="N/A",
            upd="N/A",
            settings={"fuid_elaborado_por": "JOSE TORRES", "fuid_oficina": "COMERCIAL"},
            inventariado=date(2026, 10, 26),
        )
        assert celdas[5] == "COMERCIAL"
        assert celdas[23] == "JOSE TORRES"


class TestEntregaEnCarpeta:
    """Una corrida de carpeta entrega el FUID de Kingspan junto a la UPD, en
    lugar de la planilla genérica, y arma con ellos el libro de la carpeta."""

    def test_el_fuid_llega_al_destino_y_al_libro_general(self, tmp_path):
        from resolutions.adapters.excel_inventory import SUFFIX
        from resolutions.api.folders import FolderRun, FolderRunner
        from resolutions.api.jobs import JobRegistry
        from resolutions.api.settings import Settings

        async def _no_corre(job) -> None:
            pass

        ajustes = Settings(output_dir=tmp_path / "outputs")
        registro = JobRegistry()
        runner = FolderRunner(registro, ajustes, run_job=_no_corre)
        origen = tmp_path / "origen"
        run = FolderRun(
            id="corrida", source=origen, destination=tmp_path / "destino", tipo_documento="kingspan", acta="094"
        )
        pdf = origen / "100C000999" / "UPD3869754.pdf"
        job = registro.create("100C000999/UPD3869754.pdf", pdf, tipo_documento="kingspan", acta="094")
        celdas = celdas_kingspan(
            FilaKingspan(asunto="NOTAS INTERNAS", desde="001-NI-00014161", hasta="001-NI-00014368"),
            orden=1,
            caja="100C000999",
            upd="UPD3869754",
            settings={},
            inventariado=date(2026, 10, 26),
            acta="094",
        )
        producido = ajustes.output_dir / job.id / f"UPD3869754{INVENTORY_SUFFIX}"
        escritor().write_cells([celdas], producido)
        job.report = {"outputs": [str(producido)], "inventory": {"items": []}, "notices": []}

        runner._entregar(run, job, pdf)

        entregados = {
            ruta.name for ruta in run.destination.rglob("*") if ruta.is_file() and not ruta.name.startswith(".bot-")
        }
        assert f"UPD3869754{INVENTORY_SUFFIX}" in entregados
        assert not any(nombre.endswith(SUFFIX) for nombre in entregados)

        general = inventario_general(run.destination)
        fila = next(openpyxl.load_workbook(general)[HOJA].iter_rows(min_row=PRIMERA_FILA, values_only=True))
        assert (fila[10], fila[14], fila[15], fila[25]) == ("001-NI-00014161", "100C000999", "UPD3869754", "094")

    def test_el_general_no_cuenta_lo_apartado_al_rehacer(self, tmp_path):
        celdas = celdas_kingspan(
            FilaKingspan(asunto="NOTAS INTERNAS"),
            orden=1,
            caja="100C000999",
            upd="UPD0000001",
            settings={},
            inventariado=date(2026, 10, 26),
        )
        escritor().write_cells([celdas], tmp_path / "100C000999" / f"UPD0000001{INVENTORY_SUFFIX}")
        apartado = tmp_path / "100C000999" / "_reemplazados" / "20261026" / f"UPD0000001{INVENTORY_SUFFIX}"
        escritor().write_cells([celdas], apartado)

        hoja = openpyxl.load_workbook(inventario_general(tmp_path))[HOJA]
        filas = [f for f in hoja.iter_rows(min_row=PRIMERA_FILA, values_only=True) if f[0] is not None]
        assert len(filas) == 1
