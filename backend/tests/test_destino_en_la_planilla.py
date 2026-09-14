"""La planilla dice la carpeta en la que están los PDF.

Dos fallos de la misma familia, los dos por el mismo motivo: quien escribe la
planilla no siempre sabe dónde acabará la entrega.

El primero era una ruta que se quedó fuera. La de resoluciones escribe su
planilla desde `ProcessDocument` y no desde el worker, así que cuando las otras
tres empezaron a declarar el destino, ésa siguió con la columna vacía.

El segundo es de tiempo. La carpeta definitiva lleva el NIC del expediente, y
el NIC no se conoce hasta haber leído el documento: el destino que se guarda al
crear el trabajo se calcula con el nombre del PDF de origen. La planilla que
queda junto al trabajo -- la que baja el botón de la pantalla -- decía entonces
una carpeta y los archivos estaban en otra.
"""

from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("openpyxl")
import openpyxl  # noqa: E402

from resolutions.adapters.excel_inventory import SUFFIX  # noqa: E402

INFORME = {
    "document": "UPD2365925.pdf",
    "page_count": 2,
    "nic": "5825181",
    "review_queue": [],
    "inventory": {
        "source_document": "UPD2365925.pdf",
        "source_pages": 2,
        "items": [
            {
                "code": "01",
                "title": "páginas 1-2",
                "type": "DERECHO DE PETICION",
                "fecha": "2021-12-10",
                "file_name": "UPD2365925/01_DERECHO-DE-PETICION.pdf",
                "page_count": 2,
                "first_page": 1,
                "last_page": 2,
                "page_numbers": [1, 2],
            }
        ],
    },
}


async def _no_corre(job) -> None:
    """Estas pruebas no procesan nada: miran lo que la entrega deja escrito."""


def texto_de(hoja: Path) -> str:
    libro = openpyxl.load_workbook(hoja)
    return "\n".join(
        str(celda.value)
        for nombre in libro.sheetnames
        for fila in libro[nombre].iter_rows()
        for celda in fila
        if celda.value is not None
    )


CUERPO = (
    "Por medio de la presente resolucion se resuelve la solicitud radicada en "
    "el asunto de la referencia, en los terminos que se exponen a continuacion "
    "y con base en la normatividad vigente. RESUELVE: ARTICULO PRIMERO."
)


class TestLaRutaDeResolucionesTambienDeclaraElDestino:
    """Escribía su planilla desde el caso de uso, y por eso se quedó fuera del
    arreglo. Ahora la escribe el worker, como las otras tres rutas, con el
    destino que viene en el trabajo."""

    @pytest.fixture
    def legajo(self, tmp_path):
        pymupdf = pytest.importorskip("pymupdf")
        documento = pymupdf.open()
        for encabezado in ("RESOLUCION No. 00412", ""):
            page = documento.new_page()
            if encabezado:
                page.insert_textbox(
                    pymupdf.Rect(56, 64, 540, 96),
                    encabezado,
                    fontsize=12,
                    align=pymupdf.TEXT_ALIGN_CENTER,
                )
            page.insert_textbox(pymupdf.Rect(56, 110, 540, 700), CUERPO, fontsize=11)
        ruta = tmp_path / "legajo.pdf"
        documento.save(ruta)
        documento.close()
        return ruta

    def test_la_planilla_lleva_la_carpeta_de_destino(self, legajo, tmp_path):
        from resolutions.api.settings import Settings
        from resolutions.api.worker import _split_job
        from resolutions.application.task import TaskKind

        destino = str(tmp_path / "destino" / "108C000094")
        _split_job(
            {
                "job_id": "trabajo",
                "source": str(legajo),
                "filename": "RESOLUCIONES 00412.pdf",
                "task": str(TaskKind.SPLIT),
                "destination": destino,
                "settings": Settings(output_dir=tmp_path / "outputs").as_worker_payload(),
            },
            TaskKind.SPLIT,
        )

        hoja = next(iter((tmp_path / "outputs" / "trabajo").rglob(f"*{SUFFIX}")))
        assert destino in texto_de(hoja)


class TestLaPlanillaDelTrabajoSeCorrigeAlEntregar:
    """Con la carpeta en la que los PDF acabaron de verdad."""

    @staticmethod
    def _preparar(tmp_path):
        from resolutions.api.folders import FolderRun, FolderRunner
        from resolutions.api.jobs import JobRegistry
        from resolutions.api.settings import Settings

        ajustes = Settings(output_dir=tmp_path / "outputs")
        registro = JobRegistry()
        runner = FolderRunner(registro, ajustes, run_job=_no_corre)
        run = FolderRun(
            id="corrida", source=tmp_path / "origen", destination=tmp_path / "destino"
        )
        job = registro.create("UPD2365925.pdf", tmp_path / "origen" / "UPD2365925.pdf")
        job.report = INFORME

        # La planilla que dejó el worker, con el destino que se supuso al
        # crear el trabajo: el nombre del PDF, porque el NIC no se sabía.
        carpeta_del_trabajo = ajustes.output_dir / job.id / "UPD2365925"
        carpeta_del_trabajo.mkdir(parents=True, exist_ok=True)
        from resolutions.adapters.excel_inventory import ExcelInventory

        ExcelInventory().write(
            INFORME,
            carpeta_del_trabajo,
            delivered_to=str(tmp_path / "destino" / "UPD2365925"),
        )
        return runner, run, job, carpeta_del_trabajo

    def test_antes_de_entregar_dice_la_carpeta_supuesta(self, tmp_path):
        _, _, _, carpeta = self._preparar(tmp_path)
        hoja = next(iter(carpeta.glob(f"*{SUFFIX}")))
        assert "UPD2365925" in texto_de(hoja)

    def test_al_entregar_pasa_a_decir_la_del_nic(self, tmp_path):
        runner, run, job, carpeta = self._preparar(tmp_path)
        definitiva = tmp_path / "destino" / "5825181"

        runner._corregir_planilla_del_trabajo(run, job, definitiva)

        hoja = next(iter(carpeta.glob(f"*{SUFFIX}")))
        assert str(definitiva) in texto_de(hoja)

    def test_y_no_deja_una_segunda_planilla_discrepando(self, tmp_path):
        """Se sobrescribe la que hay, en su sitio."""
        runner, run, job, carpeta = self._preparar(tmp_path)
        runner._corregir_planilla_del_trabajo(run, job, tmp_path / "destino" / "5825181")
        assert len(list(carpeta.glob(f"*{SUFFIX}"))) == 1

    def test_sin_planilla_que_corregir_no_pasa_nada(self, tmp_path):
        """Un trabajo de inventario no escribe PDF ni planilla del documento."""
        runner, run, job, carpeta = self._preparar(tmp_path)
        for hoja in carpeta.glob(f"*{SUFFIX}"):
            hoja.unlink()
        runner._corregir_planilla_del_trabajo(run, job, tmp_path / "destino" / "5825181")
        assert not list(carpeta.glob(f"*{SUFFIX}"))
