"""La planilla del inventario se escribe sola, y se baja de un clic.

El operador no tiene que pedirla: al terminar un documento queda junto a los
PDF que describe. Hasta ahora sólo la escribía la ruta de resoluciones, así que
quien separaba una caja revuelta o partía un libro de folios se quedaba con los
archivos y sin el papel que dice qué son -- que en la ruta de cajas es
especialmente grave, porque sus PDF se llaman por un número de orden y sin el
listado no dicen nada por sí solos.
"""

from __future__ import annotations

import pytest

pymupdf = pytest.importorskip("pymupdf")
pytest.importorskip("openpyxl")

from resolutions.adapters.excel_inventory import SUFFIX  # noqa: E402
from resolutions.api.settings import Settings  # noqa: E402
from resolutions.api.worker import process_document_job  # noqa: E402
from resolutions.application.task import TaskKind  # noqa: E402

CUERPO = (
    "Por medio de la presente me permito dar respuesta a la solicitud radicada "
    "en el asunto de la referencia, en los terminos que se exponen a "
    "continuacion y con base en la normatividad vigente."
)

PAGINAS = [
    ("NOTIFICACION POR AVISO", "Página 1 de 2"),
    ("", "Página 2 de 2"),
    ("FACTURA DE VENTA", "Página 1 de 1"),
]


@pytest.fixture
def caja(tmp_path):
    documento = pymupdf.open()
    for encabezado, paginacion in PAGINAS:
        page = documento.new_page()
        if encabezado:
            page.insert_textbox(
                pymupdf.Rect(56, 64, 540, 96),
                encabezado,
                fontsize=12,
                align=pymupdf.TEXT_ALIGN_CENTER,
            )
        page.insert_textbox(pymupdf.Rect(56, 110, 540, 700), CUERPO, fontsize=11)
        page.insert_textbox(pymupdf.Rect(56, 720, 540, 760), paginacion, fontsize=9)
    ruta = tmp_path / "caja.pdf"
    documento.save(ruta)
    documento.close()
    return ruta


@pytest.fixture
def procesada(caja, tmp_path):
    informe = process_document_job(
        {
            "job_id": "trabajo",
            "source": str(caja),
            "filename": "UPD2365925.pdf",
            "task": str(TaskKind.SEGMENT),
            "operator": "Eduver Andrés",
            "settings": Settings(output_dir=tmp_path / "outputs").as_worker_payload(),
        }
    )
    return informe, tmp_path / "outputs" / "trabajo"


class TestLaPlanillaSeEscribeSola:
    def test_queda_una_planilla_en_el_trabajo(self, procesada):
        _, salida = procesada
        assert list(salida.rglob(f"*{SUFFIX}"))

    def test_junto_a_los_pdf_que_describe(self, procesada):
        """En la carpeta de la caja, no en la raíz del trabajo: quien abre el
        disco encuentra el listado al lado de los archivos que lista."""
        _, salida = procesada
        hoja = next(iter(salida.rglob(f"*{SUFFIX}")))
        assert hoja.parent.name == "UPD2365925"
        assert list(hoja.parent.glob("*.pdf"))

    def test_se_llama_como_el_documento_de_origen(self, procesada):
        _, salida = procesada
        hoja = next(iter(salida.rglob(f"*{SUFFIX}")))
        assert hoja.name.startswith("UPD2365925")

    def test_y_lista_lo_que_de_verdad_salio(self, procesada):
        import openpyxl

        informe, salida = procesada
        hoja = next(iter(salida.rglob(f"*{SUFFIX}")))
        libro = openpyxl.load_workbook(hoja)
        texto = "\n".join(
            str(celda.value)
            for fila in libro["Resoluciones"].iter_rows()
            for celda in fila
            if celda.value is not None
        )
        for item in informe["inventory"]["items"]:
            assert item["file_name"].rsplit("/", 1)[-1] in texto

    def test_el_tipo_documental_esta_en_la_planilla(self, procesada):
        import openpyxl

        _, salida = procesada
        hoja = next(iter(salida.rglob(f"*{SUFFIX}")))
        libro = openpyxl.load_workbook(hoja)
        texto = "\n".join(
            str(celda.value)
            for fila in libro["Resoluciones"].iter_rows()
            for celda in fila
            if celda.value is not None
        )
        assert "NOTIFICACION POR AVISO" in texto

    def test_un_fallo_de_la_planilla_no_se_lleva_el_trabajo(self, caja, tmp_path, monkeypatch):
        """Los PDF ya están escritos cuando esto corre."""
        from resolutions.adapters import excel_inventory

        class Rota:
            def write(self, *_a, **_k):
                raise OSError("el disco dijo que no")

        monkeypatch.setattr(excel_inventory, "ExcelInventory", Rota)
        informe = process_document_job(
            {
                "job_id": "otro",
                "source": str(caja),
                "filename": "caja.pdf",
                "task": str(TaskKind.SEGMENT),
                "settings": Settings(output_dir=tmp_path / "outputs").as_worker_payload(),
            }
        )
        assert informe["outputs"]


class TestSeBajaDeUnClic:
    def test_el_endpoint_la_encuentra_dentro_de_la_carpeta_de_la_caja(
        self, client, tmp_path, monkeypatch
    ):
        """Buscándola sólo en la raíz del trabajo contestaba 404 justo en la
        ruta que más la necesita."""
        pytest.importorskip("httpx")
        from resolutions.adapters.excel_inventory import SUFFIX as sufijo
        from resolutions.api import main

        carpeta = main.settings.output_dir / "trabajo" / "UPD2365925"
        carpeta.mkdir(parents=True, exist_ok=True)
        (carpeta / f"UPD2365925{sufijo}").write_bytes(b"PK\x03\x04 planilla")

        respuesta = client.get("/api/jobs/trabajo/inventory.xlsx")
        assert respuesta.status_code == 200

    def test_sin_planilla_lo_dice_en_vez_de_reventar(self, client):
        pytest.importorskip("httpx")
        assert client.get("/api/jobs/fantasma/inventory.xlsx").status_code == 404
